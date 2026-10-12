"""Cabecalho e rodape (M5): .docx, oraculo, validador, limites, tela, edicao no lugar e comandos.

- ida e volta: padrao, primeira pagina (titlePg) e pares (evenAndOddHeaders), com imagem
  (relacoes proprias da parte), tabela e campos; distancias do pgMar; idempotencia;
- python-docx como oraculo: o Word ve as mesmas partes, opcoes e textos;
- validador: parte sem tipo de conteudo, cabecalho sem paragrafo, sectPr fora de ordem e
  r:id sem relacao no cabecalho sao acusados;
- leitura como o Word grava: heranca entre secoes, secoes com cabecalhos diferentes,
  referencia quebrada, parte grande demais, DOCTYPE no cabecalho, numeracao estranha;
- tela: cada pagina pinta a sua variante, com o numero certo; esmaecido fora da edicao;
- edicao no lugar com mouse e teclado reais: duplo clique entra, digita, Ctrl+Z e Ctrl+Y
  dentro do cabecalho, Esc e duplo clique no corpo saem; o corpo desce quando o
  cabecalho cresce;
- faixa: Numero de Pagina, Pagina X de Y, primeira pagina, pares e impares, formato,
  numero inicial, remover, fechar; salvar e reabrir; amostra para o Word.
"""

from __future__ import annotations

import io
import pathlib

from ajudantes import (RAIZ, checa, checa_igual, drenar_eventos, pasta_temporaria, preparar_qt,
                       pular, resumir, secao, usar_fonte_windows)
import fixtures_docx as fx

QT = preparar_qt()

from editdoc.formato_docx import conferencia, ordem  # noqa: E402
from editdoc.formato_docx.erros import ErroDocx  # noqa: E402
from editdoc.formato_docx.escritor import escrever, escrever_com_avisos  # noqa: E402
from editdoc.formato_docx.leitor import ler  # noqa: E402
from editdoc.modelo.documento import Documento, Paragrafo, Run, normalizar, paragrafos_de  # noqa: E402
from editdoc.modelo.pagina import Pagina  # noqa: E402

PASTA_AMOSTRAS = RAIZ / "tests" / "saida"


def codigos(avisos) -> dict[str, int]:
    return {a.codigo: a.contagem for a in avisos}


def texto_de(blocos) -> str:
    return " | ".join("".join(r.texto for r in p.runs) for p in paragrafos_de(blocos))


def faixas_iguais(a: Documento, b: Documento) -> bool:
    na, nb = normalizar(a), normalizar(b)
    return na.cabecalhos == nb.cabecalhos and na.rodapes == nb.rodapes


# -- .docx ----------------------------------------------------------------------------

def testar_ida_e_volta() -> None:
    secao(".docx: padrao, primeira, par, imagem, tabela e campos; ida e volta e idempotencia")
    doc = fx.documento_com_cabecalhos(20)
    dados, avisos = escrever_com_avisos(doc)
    checa_igual(avisos, [], "gravar cabecalho e rodape nao gera aviso")
    partes = fx.partes_de(dados)
    checa_igual(sorted(n for n in partes if n.startswith(("word/header", "word/footer"))),
                ["word/footer1.xml", "word/footer2.xml", "word/header1.xml", "word/header2.xml",
                 "word/header3.xml"], "uma parte por variante: 3 cabecalhos e 2 rodapes")
    rels_cab = partes.get("word/_rels/header1.xml.rels", b"").decode("utf-8")
    checa("media/image1.png" in rels_cab, "a imagem do cabecalho tem relacao na propria parte")
    ct = partes["[Content_Types].xml"].decode("utf-8")
    checa(ct.count("wordprocessingml.header+xml") == 3 and ct.count("wordprocessingml.footer+xml") == 2,
          "cada parte tem o seu tipo de conteudo (Override)")
    xml = partes["word/document.xml"].decode("utf-8")
    checa('w:type="first"' in xml and 'w:type="even"' in xml and 'w:type="default"' in xml
          and "<w:titlePg/>" in xml, "sectPr com as tres referencias e titlePg")
    checa("<w:evenAndOddHeaders/>" in partes["word/settings.xml"].decode("utf-8"),
          "settings.xml com evenAndOddHeaders")
    checa('w:header="708"' in xml and 'w:footer="708"' in xml, "distancias do cabecalho e do rodape no pgMar")
    checa_igual(ordem.validar(partes), [], "o pacote passa no validador")
    lido, avisos_l = ler(dados)
    checa_igual(avisos_l, [], "o proprio arquivo abre sem aviso")
    checa(faixas_iguais(doc, lido), "cabecalhos e rodapes voltam iguais")
    checa(normalizar(lido).paragrafos == normalizar(doc).paragrafos, "o corpo volta igual")
    checa_igual((lido.pagina.primeira_diferente, lido.pagina.pares_impares,
                 lido.pagina.cabecalho_tw, lido.pagina.rodape_tw), (True, True, 708, 708),
                "opcoes e distancias voltam")
    checa(escrever(lido) == escrever(ler(escrever(lido))[0]), "a segunda ida produz os mesmos bytes")
    conferencia.conferir(dados, doc)
    checa(True, "a conferencia aceita o arquivo com cabecalhos")
    errado = normalizar(doc)
    errado.cabecalhos["primeira"][0].runs[0].texto = "outro texto"
    try:
        conferencia.conferir(dados, errado)
        checa(False, "contraprova: cabecalho diferente passaria pela conferencia")
    except conferencia.FalhaDeConferencia as erro:
        checa("cabecalho" in str(erro), f"contraprova: cabecalho diferente e' acusado ({erro})")
    sem_rodape = normalizar(doc)
    sem_rodape.rodapes.pop("par")
    try:
        conferencia.conferir(dados, sem_rodape)
        checa(False, "contraprova: rodape a mais passaria pela conferencia")
    except conferencia.FalhaDeConferencia:
        checa(True, "contraprova: variante de rodape a mais e' acusada")


def testar_oraculo() -> None:
    secao("python-docx como oraculo: partes, opcoes e textos que o Word ve")
    import docx
    d = docx.Document(io.BytesIO(escrever(fx.documento_com_cabecalhos(5))))
    s = d.sections[0]
    checa(s.different_first_page_header_footer, "primeira pagina diferente (titlePg)")
    checa(d.settings.odd_and_even_pages_header_footer, "pares e impares diferentes (settings)")
    checa(not s.header.is_linked_to_previous and not s.footer.is_linked_to_previous,
          "cabecalho e rodape padrao existem (nao herdados)")
    checa("Cabeçalho das páginas ímpares" in s.header.paragraphs[0].text, "texto do cabecalho padrao")
    checa_igual(s.first_page_header.paragraphs[0].text, "Capa: sem número de página",
                "texto do cabecalho da primeira pagina")
    checa_igual(len(s.even_page_header.tables), 1, "a tabela do cabecalho par")
    checa_igual(s.footer.paragraphs[0].text, "Página 1 de 4", "rodape com o resultado dos campos")
    checa_igual(s.header_distance.twips if hasattr(s.header_distance, "twips") else round(s.header_distance / 635),
                708, "distancia do cabecalho")
    imagens = [r for r in s.header.part.rels.values() if r.reltype.endswith("/image")]
    checa_igual(len(imagens), 1, "a imagem do cabecalho e' alcancada pela parte do cabecalho")


def testar_validador() -> None:
    secao("validador: tipo de conteudo, cabecalho vazio, ordem do sectPr e r:id do cabecalho")
    dados = escrever(fx.documento_com_cabecalhos(2))
    partes = fx.partes_de(dados)
    checa_igual(ordem.validar(partes), [], "o arquivo bom passa")
    sab = dict(partes)
    ct = sab["[Content_Types].xml"].decode("utf-8")
    sab["[Content_Types].xml"] = ct.replace('<Override PartName="/word/footer2.xml"', '<Override PartName="/word/x.xml"').encode("utf-8")
    checa(any("footer2.xml" in p for p in ordem.validar(sab)), "contraprova: rodape sem tipo de conteudo")
    sab = dict(partes)
    sab["word/header2.xml"] = fx.parte_faixa("")
    checa(any("nao termina em paragrafo" in p for p in ordem.validar(sab)),
          "contraprova: cabecalho sem paragrafo")
    sab = dict(partes)
    xml = sab["word/document.xml"].decode("utf-8")
    sab["word/document.xml"] = xml.replace("<w:titlePg/>", "").replace("<w:pgSz", "<w:titlePg/><w:pgSz", 1).encode("utf-8")
    checa(any("sectPr" in p for p in ordem.validar(sab)), "contraprova: titlePg antes de pgSz")
    sab = dict(partes)
    sab["word/_rels/header1.xml.rels"] = sab["word/_rels/header1.xml.rels"].replace(b'Id="rId1"', b'Id="rId9"')
    checa(any("header1.xml" in p and "r:id" in p for p in ordem.validar(sab)),
          "contraprova: imagem do cabecalho sem relacao na parte")


def _referencia(tipo: str, rid: str, onde: str = "header") -> str:
    return f'<w:{onde}Reference w:type="{tipo}" r:id="{rid}"/>'


def testar_leitura_do_word() -> None:
    secao("leitura como o Word grava: heranca, secoes, referencia quebrada, limites e numeracao")
    cab = fx.parte_faixa('<w:p><w:pPr><w:jc w:val="right"/></w:pPr><w:r><w:t>Relatório</w:t></w:r>'
                         '<w:fldSimple w:instr=" PAGE "><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>')
    rod = fx.parte_faixa("<w:p>" + fx.campo_complexo("NUMPAGES", "3") + "</w:p>", "ftr")
    base = dict(partes_extra={"word/header7.xml": cab, "word/footer7.xml": rod},
                relacoes_extra=[("rId70", fx.RT_HEADER, "header7.xml"), ("rId71", fx.RT_FOOTER, "footer7.xml")],
                tipos_extra=[("/word/header7.xml", fx.CT_HEADER), ("/word/footer7.xml", fx.CT_FOOTER)])
    dados = fx.docx_com_xml("<w:p><w:r><w:t>corpo</w:t></w:r></w:p>",
                            _referencia("default", "rId70") + _referencia("default", "rId71", "footer")
                            .replace("footerReference", "footerReference"), **base)
    modelo, avisos = ler(dados)
    checa_igual(avisos, [], "cabecalho e rodape do Word abrem sem aviso")
    cab_runs = modelo.cabecalhos["padrao"][0].runs
    checa_igual([(r.texto, r.campo) for r in cab_runs], [("Relatório", None), ("1", "PAGE")],
                "cabecalho com texto e campo simples")
    checa_igual(modelo.cabecalhos["padrao"][0].alinhamento, "direita", "alinhamento do cabecalho")
    checa_igual(modelo.rodapes["padrao"][0].runs[0].campo, "NUMPAGES", "rodape com campo complexo")

    # Heranca: a primeira secao declara o cabecalho, a ultima nao declara nada.
    corpo = ("<w:p><w:pPr><w:sectPr>" + _referencia("default", "rId70")
             + "</w:sectPr></w:pPr><w:r><w:t>secao 1</w:t></w:r></w:p><w:p><w:r><w:t>secao 2</w:t></w:r></w:p>")
    modelo, avisos = ler(fx.docx_com_xml(corpo, **base))
    checa("padrao" in modelo.cabecalhos and "CABECALHO_SECAO" not in codigos(avisos),
          "secao sem referencia herda o cabecalho da anterior, sem aviso de cabecalho")
    # Duas secoes com cabecalhos diferentes.
    outro = fx.parte_faixa("<w:p><w:r><w:t>Outro</w:t></w:r></w:p>")
    base2 = dict(base)
    base2["partes_extra"] = {**base["partes_extra"], "word/header8.xml": outro}
    base2["relacoes_extra"] = base["relacoes_extra"] + [("rId80", fx.RT_HEADER, "header8.xml")]
    base2["tipos_extra"] = base["tipos_extra"] + [("/word/header8.xml", fx.CT_HEADER)]
    modelo, avisos = ler(fx.docx_com_xml(corpo, _referencia("default", "rId80"), **base2))
    checa("CABECALHO_SECAO" in codigos(avisos), "secoes com cabecalhos diferentes: aviso CABECALHO_SECAO")
    checa_igual(texto_de(modelo.cabecalhos["padrao"]), "Outro", "fica o cabecalho da ultima secao")

    # Arquivo hostil: 400 secoes, cada uma com a sua referencia. So' a ultima e' lida.
    import time
    muitas = dict(base)
    muitas["relacoes_extra"] = base["relacoes_extra"] + [
        (f"rIdM{i}", fx.RT_HEADER, "header7.xml") for i in range(400)]
    corpo_muitas = "".join(f"<w:p><w:pPr><w:sectPr>{_referencia('default', f'rIdM{i}')}"
                           f"{_referencia('first', f'rIdM{i}')}</w:sectPr></w:pPr></w:p>"
                           for i in range(400))
    inicio = time.perf_counter()
    modelo, avisos = ler(fx.docx_com_xml(corpo_muitas, _referencia("default", "rId70"), **muitas))
    gasto = time.perf_counter() - inicio
    checa(gasto < 2.0 and "CABECALHO_SECAO" in codigos(avisos),
          f"400 secoes com cabecalho: le' em {gasto:.2f} s (limite 2 s), com aviso")
    checa_igual(sorted(modelo.cabecalhos), ["padrao", "primeira"],
                "so' as variantes da ultima secao (com a heranca) sao lidas")

    modelo, avisos = ler(fx.docx_com_xml("<w:p/>", _referencia("default", "rId99"), **base))
    checa(codigos(avisos).get("CABECALHO_RODAPE") == 1 and not modelo.cabecalhos,
          "referencia sem relacao: aviso CABECALHO_RODAPE")
    modelo, avisos = ler(fx.docx_com_xml("<w:p/>", _referencia("default", "rId71"), **base))
    checa("CABECALHO_RODAPE" in codigos(avisos), "cabecalho que aponta para rodape: aviso")
    modelo, avisos = ler(fx.docx_com_xml("<w:p/>", _referencia("default", "rId70"), **base),
                         limites={"cabecalho_bytes": 100})
    checa(codigos(avisos).get("CABECALHO_GRANDE") == 1 and not modelo.cabecalhos,
          "cabecalho acima do limite de tamanho: nao e' lido, com aviso")
    hostil = dict(base)
    hostil["partes_extra"] = {**base["partes_extra"], "word/header7.xml":
                              b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "b">]>' + cab[cab.index(b"<w:hdr"):]}
    try:
        ler(fx.docx_com_xml("<w:p/>", _referencia("default", "rId70"), **hostil))
        checa(False, "DOCTYPE no cabecalho deveria ser recusado")
    except ErroDocx as erro:
        checa_igual(erro.codigo, "XML_DOCTYPE_PROIBIDO", "DOCTYPE no cabecalho e' recusado")

    sect = '<w:pgNumType w:fmt="chicago" w:start="-3"/>'
    dados = fx.docx_com_xml("<w:p/>")
    partes = fx.partes_de(dados)
    xml = partes["word/document.xml"].decode("utf-8").replace("</w:sectPr>", sect + "<w:titlePg w:val=\"0\"/></w:sectPr>")
    partes["word/document.xml"] = xml.encode("utf-8")
    modelo, avisos = ler(fx.montar(partes))
    checa_igual(codigos(avisos), {"NUMERO_PAGINA": 2}, "formato chicago e inicio negativo: aviso NUMERO_PAGINA")
    checa_igual((modelo.pagina.formato_numero, modelo.pagina.numero_inicial, modelo.pagina.primeira_diferente),
                ("decimal", None, False), "numeracao estranha cai para 1, 2, 3; titlePg val=0 desliga")
    modelo, _ = ler(fx.docx_com_xml("<w:p/>", settings_extra='<w:evenAndOddHeaders w:val="1"/>'))
    checa(modelo.pagina.pares_impares, "evenAndOddHeaders do settings.xml e' lido")


# -- tela e edicao ------------------------------------------------------------------------

def _janela(modelo: Documento, largura: int = 1100, altura: int = 1400):
    from editdoc import configuracao
    from editdoc.interface.janela_principal import JanelaPrincipal
    janela = JanelaPrincipal(configuracao.padrao())
    janela.resize(largura, altura)
    janela.show()
    janela.carregar(modelo, [], pathlib.Path("teste_m5.docx"))
    drenar_eventos()
    return janela


def _ponto(painel, n: int, x_tw: float, y_tw: float):
    from PySide6.QtCore import QPoint
    m = painel.mapa()
    x, y, _w, _h = m.folha_na_tela(n)
    return QPoint(round(x + x_tw * m.escala), round(y + y_tw * m.escala))


def _fechar(janela) -> None:
    janela.painel.documento_corpo().setModified(False)
    janela.close()
    drenar_eventos()


def testar_paginas_na_tela() -> None:
    secao("tela: cada pagina com a sua variante e o numero certo; esmaecido fora da edicao")
    from editdoc.interface.vista_paginas import rotulo_da_faixa
    from editdoc.ponte_qt import campos

    vistos: list[tuple] = []
    original = campos.ManipuladorDeCampo.drawObject

    def espiao(self, painter, rect, doc, pos, fmt):
        instrucao = campos.instrucao_de(fmt)
        indice = self._indice(rect.center().y())
        vistos.append((self.contexto.indice_fixo, instrucao, self.contexto.valor(instrucao, indice)))
        return original(self, painter, rect, doc, pos, fmt)

    campos.ManipuladorDeCampo.drawObject = espiao
    try:
        janela = _janela(fx.documento_com_cabecalhos(), 1100, 1500)
        painel = janela.painel
        painel.definir_zoom(0.25)
        drenar_eventos()
        total = painel.total_de_paginas()
        checa(total >= 4, f"a amostra tem pelo menos 4 paginas ({total})")
        vistos.clear()
        painel._limpar_cache()  # o corpo vem dos tiles: sem limpar, o campo nao e' repintado
        painel.viewport().grab()
        rodape = {}
        for indice, instrucao, valor in vistos:
            if indice is not None:
                rodape.setdefault(indice, set()).add((instrucao, valor))
        checa(0 not in rodape, "pagina 1 (primeira diferente, sem rodape proprio): nenhum numero")
        checa_igual(rodape.get(1), {("PAGE", "2")}, "pagina 2 (par): rodape par com o numero 2")
        checa_igual(rodape.get(2), {("PAGE", "3"), ("NUMPAGES", str(total))},
                    f"pagina 3 (impar): 'Página 3 de {total}'")
        checa_igual(rodape.get(3), {("PAGE", "4")}, "pagina 4 (par): numero 4")
        corpo = {v for i, ins, v in vistos if i is None and ins == "PAGE"}
        checa("1" in corpo and str(total) in corpo, "os campos do corpo mostram a pagina onde estao")
        pagina = painel.configuracao_pagina()
        checa_igual([rotulo_da_faixa("cabecalho", i, pagina) for i in range(3)],
                    ["Cabeçalho - Primeira página", "Cabeçalho - Página par", "Cabeçalho - Página ímpar"],
                    "rotulos das variantes")
        _fechar(janela)
    finally:
        campos.ManipuladorDeCampo.drawObject = original

    # Cinza claro fora da edicao, escuro na edicao (pixels do cabecalho de "M" pretos).
    doc = Documento(paragrafos=[Paragrafo(runs=[Run("corpo")])],
                    cabecalhos={"padrao": [Paragrafo(runs=[Run("MMMMMMMMMMMM", tamanho_meio_pt=40)])]})
    janela = _janela(doc, 1000, 900)
    painel = janela.painel
    m = painel.mapa()
    x, y, _w, _h = m.folha_na_tela(0)
    pagina = painel.configuracao_pagina()
    area = (round(x + pagina.margem_esq_tw * m.escala), round(y + pagina.cabecalho_tw * m.escala),
            round(3000 * m.escala), round(500 * m.escala))

    def mais_escuro() -> int:
        img = painel.viewport().grab().toImage()
        menor = 255
        for i in range(area[0], area[0] + area[2], 2):
            for j in range(area[1], area[1] + area[3], 2):
                c = img.pixelColor(i, j)
                menor = min(menor, (c.red() + c.green() + c.blue()) // 3)
        return menor

    fora = mais_escuro()
    painel.entrar_na_faixa("cabecalho", 0)
    drenar_eventos()
    dentro = mais_escuro()
    checa(fora > 90, f"fora da edicao o cabecalho e' cinza claro (pixel mais escuro {fora})")
    checa(dentro < 60, f"na edicao o cabecalho fica escuro (pixel mais escuro {dentro})")
    painel.sair_da_faixa()
    _fechar(janela)


def testar_edicao_no_lugar() -> None:
    secao("edicao no lugar: duplo clique, digitar, Ctrl+Z/Ctrl+Y no cabecalho, Esc, corpo desce")
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest

    doc = Documento(paragrafos=[Paragrafo(runs=[Run(f"linha {i} do corpo do documento")]) for i in range(80)],
                    cabecalhos={"padrao": [Paragrafo(runs=[Run("Topo")])]})
    janela = _janela(doc)
    painel = janela.painel
    corpo_antes = painel.toPlainText()
    paginas_antes = painel.total_de_paginas()
    vp = painel.viewport()
    QTest.mouseDClick(vp, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier,
                      _ponto(painel, 0, 3000, 900))
    drenar_eventos()
    faixa = painel.editando_faixa()
    checa(faixa is not None and faixa.onde == "cabecalho", "duplo clique no alto da pagina edita o cabecalho")
    checa_igual(janela.faixa.aba_atual(), "Cabeçalho e Rodapé", "a aba contextual aparece e fica ativa")
    checa(janela.comandos.acao("fechar_cabecalho").isEnabled(), "Fechar cabeçalho e rodapé habilitado")
    QTest.keyClick(vp, Qt.Key.Key_End)
    QTest.keyClicks(vp, " ABC")
    checa_igual(faixa.doc.toPlainText(), "Topo ABC", "o texto digitado entra no cabecalho")
    checa_igual(painel.toPlainText(), corpo_antes, "o corpo nao muda")
    checa(janela.sujo() and janela.windowTitle().startswith("*"), "editar o cabecalho deixa o documento alterado")
    checa(not painel.documento_corpo().isUndoAvailable(), "o desfazer do corpo continua vazio")
    QTest.keyClick(vp, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
    checa_igual(faixa.doc.toPlainText(), "Topo", "Ctrl+Z desfaz a digitacao no cabecalho")
    QTest.keyClick(vp, Qt.Key.Key_Y, Qt.KeyboardModifier.ControlModifier)
    checa_igual(faixa.doc.toPlainText(), "Topo ABC", "Ctrl+Y refaz no cabecalho")
    QTest.keyClick(vp, Qt.Key.Key_B, Qt.KeyboardModifier.ControlModifier)
    QTest.keyClicks(vp, "N")
    checa(any(r.negrito for p in janela.painel.faixas_do_modelo()[0]["padrao"] for r in p.runs),
          "Ctrl+B formata o texto do cabecalho")

    # O corpo desce quando o cabecalho passa da margem.
    for _ in range(14):
        QTest.keyClick(vp, Qt.Key.Key_Return)
    drenar_eventos()
    efetiva = painel.layout_pag.pagina
    checa(efetiva.margem_sup_tw > 1440, f"cabecalho alto: o corpo desce (margem efetiva {efetiva.margem_sup_tw} tw)")
    checa(faixa.altura() + painel.configuracao_pagina().cabecalho_tw <= efetiva.margem_sup_tw + 1,
          "o cabecalho nao sobrepoe o corpo")
    checa(painel.total_de_paginas() > paginas_antes, "o corpo encolheu: mais paginas")
    checa_igual(painel.configuracao_pagina().margem_sup_tw, 1440, "a margem do arquivo nao muda")
    for _ in range(16):
        QTest.keyClick(vp, Qt.Key.Key_Z, Qt.KeyboardModifier.ControlModifier)
    drenar_eventos()
    checa_igual(painel.layout_pag.pagina.margem_sup_tw, 1440, "desfazer as linhas devolve a margem")

    QTest.keyClick(vp, Qt.Key.Key_Escape)
    drenar_eventos()
    checa(painel.editando_faixa() is None, "Esc volta ao corpo")
    checa("Cabeçalho e Rodapé" not in janela.faixa.abas_visiveis(), "a aba contextual some")
    QTest.keyClicks(vp, "Z")
    checa(painel.toPlainText().startswith("Z"), "depois do Esc o teclado volta ao corpo, no lugar de antes")

    # Rodape pela margem de baixo; clique no corpo nao sai; duplo clique no corpo sai.
    pagina = painel.configuracao_pagina()
    QTest.mouseDClick(vp, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier,
                      _ponto(painel, 0, 3000, pagina.altura_tw - 900))
    checa(painel.editando_faixa() is not None and painel.editando_faixa().onde == "rodape",
          "duplo clique no pe' da pagina edita o rodape")
    QTest.mouseClick(vp, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier,
                     _ponto(painel, 0, 3000, 4000))
    checa(painel.editando_faixa() is not None, "clique simples no corpo nao sai da edicao (como no Word)")
    QTest.mouseClick(vp, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier,
                     _ponto(painel, 0, 3000, 900))
    checa(painel.editando_faixa().onde == "cabecalho", "clique no cabecalho passa do rodape para o cabecalho")
    QTest.mouseDClick(vp, Qt.MouseButton.LeftButton, Qt.KeyboardModifier.NoModifier,
                      _ponto(painel, 0, 3000, 4000))
    checa(painel.editando_faixa() is None, "duplo clique no corpo volta ao corpo")
    _fechar(janela)


def testar_comandos_e_salvar(tmp: pathlib.Path) -> None:
    secao("faixa: Número de Página, Página X de Y, opcoes, formato, inicio, remover; salvar e reabrir")
    from PySide6.QtWidgets import QApplication

    from editdoc.interface import dialogos_modais
    from editdoc.interface.vista_paginas import rotulo_da_faixa

    doc = Documento(paragrafos=[Paragrafo(runs=[Run(f"paragrafo {i}")]) for i in range(120)])
    janela = _janela(doc)
    painel = janela.painel
    c = janela.comandos
    c.acao("numero_x_de_y").trigger()
    drenar_eventos()
    rodape = painel.layout_pag.faixas.obter("rodape", "padrao")
    checa(rodape is not None and len(rodape.doc.toPlainText()) == len("Página ") + 1 + len(" de ") + 1,
          "Página X de Y entra no rodape")
    _cab, rod = painel.faixas_do_modelo()
    checa_igual([r.campo for r in rod["padrao"][0].runs], [None, "PAGE", None, "NUMPAGES"],
                "com os campos PAGE e NUMPAGES")
    checa_igual(rod["padrao"][0].alinhamento, "centro", "centralizado")
    checa(painel.editando_faixa() is rodape, "depois de inserir, o rodape fica em edicao")
    c.acao("numero_topo_direita").trigger()
    cab, _rod = painel.faixas_do_modelo()
    checa(cab["padrao"][0].alinhamento == "direita" and cab["padrao"][0].runs[0].campo == "PAGE",
          "numero no cabecalho, a direita")
    c.acao("numero_topo_direita").trigger()
    cab, _rod = painel.faixas_do_modelo()
    checa_igual(len(cab["padrao"]), 2, "o segundo numero vai num paragrafo novo, sem apagar o primeiro")
    painel.undo()
    cab, _rod = painel.faixas_do_modelo()
    checa_igual(len(cab["padrao"]), 1, "um desfazer tira o numero inserido")

    painel.entrar_na_faixa("cabecalho", 0)
    c.acao("primeira_diferente").trigger()
    drenar_eventos()
    checa(painel.configuracao_pagina().primeira_diferente and c.acao("primeira_diferente").isChecked(),
          "Primeira página diferente liga a opcao e marca o botao")
    checa_igual(painel.editando_faixa().tipo, "primeira", "na pagina 1, a edicao passa ao cabecalho da primeira")
    checa_igual(rotulo_da_faixa("cabecalho", 0, painel.configuracao_pagina()), "Cabeçalho - Primeira página",
                "o rotulo diz Primeira página")
    QApplication.processEvents()
    from PySide6.QtTest import QTest
    QTest.keyClicks(painel.viewport(), "Capa")
    c.acao("pares_impares").trigger()
    checa(painel.configuracao_pagina().pares_impares, "Páginas pares e ímpares diferentes liga a opcao")
    c.acao("formato_romano").trigger()
    checa_igual(painel.configuracao_pagina().formato_numero, "lowerRoman", "formato i, ii, iii")
    original = dialogos_modais.pedir_numero_inicial
    try:
        dialogos_modais.pedir_numero_inicial = lambda _p, _a: (True, 5)
        c.acao("numero_inicial").trigger()
    finally:
        dialogos_modais.pedir_numero_inicial = original
    checa_igual(painel.configuracao_pagina().numero_inicial, 5, "Começar a numeração em 5")
    c.acao("fechar_cabecalho").trigger()
    drenar_eventos()
    checa(painel.editando_faixa() is None and not c.acao("fechar_cabecalho").isEnabled(),
          "Fechar cabeçalho e rodapé volta ao corpo")

    destino = tmp / "com_cabecalho.docx"
    original = dialogos_modais.pedir_salvar_arquivo
    try:
        dialogos_modais.pedir_salvar_arquivo = lambda *_a: str(destino)
        checa(janela.salvar_como(), "salvar o documento com cabecalho e rodape")
    finally:
        dialogos_modais.pedir_salvar_arquivo = original
    modelo, avisos = ler(destino.read_bytes())
    checa_igual(avisos, [], "o arquivo salvo abre sem aviso")
    checa_igual(ordem.validar(fx.partes_de(destino.read_bytes())), [], "o arquivo passa no validador")
    checa_igual(sorted(modelo.cabecalhos), ["padrao", "primeira"], "cabecalhos padrao e primeira no arquivo")
    checa_igual(texto_de(modelo.cabecalhos["primeira"]), "Capa", "o texto da capa")
    checa_igual([r.campo for r in modelo.rodapes["padrao"][0].runs], [None, "PAGE", None, "NUMPAGES"],
                "o rodape continua com campos (nao virou texto)")
    checa_igual((modelo.pagina.primeira_diferente, modelo.pagina.pares_impares,
                 modelo.pagina.formato_numero, modelo.pagina.numero_inicial),
                (True, True, "lowerRoman", 5), "opcoes e numeracao no arquivo")
    checa("par" not in modelo.cabecalhos, "a variante par criada e deixada vazia nao vai para o arquivo")

    c.acao("remover_rodape").trigger()
    _cab, rod = painel.faixas_do_modelo()
    checa("padrao" not in rod or texto_de(rod["padrao"]) == "",
          "Remover rodapé esvazia o rodape (e o rodape novo vazio nao vai para o arquivo)")
    _fechar(janela)

    janela = _janela(modelo)
    faixa = janela.painel.layout_pag.faixas.obter("rodape", "padrao")
    checa(faixa is not None and len(faixa.doc.toPlainText()) == 13, "reabrir: o rodape volta com os campos")
    _fechar(janela)


def testar_amostra(tmp: pathlib.Path) -> None:
    secao("amostra para abrir no Word, gravada pela propria janela")
    from editdoc.interface import dialogos_modais

    janela = _janela(fx.documento_com_cabecalhos())
    janela.painel.documento_corpo().setModified(True)
    PASTA_AMOSTRAS.mkdir(parents=True, exist_ok=True)
    amostra = PASTA_AMOSTRAS / "amostra_m5_cabecalho_rodape.docx"
    original = dialogos_modais.pedir_salvar_arquivo
    try:
        dialogos_modais.pedir_salvar_arquivo = lambda *_a: str(amostra)
        checa(janela.salvar_como(), "a janela grava a amostra")
    finally:
        dialogos_modais.pedir_salvar_arquivo = original
    dados = amostra.read_bytes()
    lido, avisos = ler(dados)
    checa(ordem.validar(fx.partes_de(dados)) == [] and avisos == [],
          f"amostra para abrir no Word: {amostra.relative_to(RAIZ)}")
    checa(faixas_iguais(fx.documento_com_cabecalhos(), lido),
          "a amostra tem os cabecalhos e rodapes do modelo")
    _fechar(janela)


def main() -> int:
    testar_ida_e_volta()
    testar_oraculo()
    testar_validador()
    testar_leitura_do_word()
    if not QT:
        return pular("PySide6 nao esta instalado neste venv")
    if not usar_fonte_windows(9):
        return pular("Segoe UI nao esta instalada")
    testar_paginas_na_tela()
    testar_edicao_no_lugar()
    with pasta_temporaria("editdoc-teste-cabecalho-") as tmp:
        testar_comandos_e_salvar(tmp)
        testar_amostra(tmp)
    return resumir()


if __name__ == "__main__":
    raise SystemExit(main())
