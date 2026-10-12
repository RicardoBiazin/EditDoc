"""Tabelas (M4): modelo, .docx, ponte Qt, visao em paginas, edicao e limites.

- ida e volta modelo -> .docx -> modelo, com mescla horizontal e vertical, bordas,
  sombreamento, cabecalho, altura de linha e tabela aninhada; idempotente;
- leitura de tabela como o Word grava (estilo "Table Grid", hMerge antigo, linha curta);
- python-docx como oraculo: o Word veria as mesmas celulas e as mesmas mesclas;
- validador do pacote acusa celula sem paragrafo, linha torta e corpo terminado em tabela;
- limites contra arquivo hostil: linhas, colunas, celulas e aninhamento;
- ponte Qt: mesclas viram rowSpan/columnSpan; bloco estrutural nao volta;
- tela: a tabela aparece, as bordas tem tinta, e a tabela atravessa paginas sem sumir;
- edicao com teclas reais: digitar, Tab, Shift+Tab, Tab na ultima celula; comandos de
  linha, coluna, mescla, divisao e exclusao, cada um um passo de desfazer.
"""

from __future__ import annotations

import io
import pathlib
import time

from ajudantes import (RAIZ, checa, checa_igual, drenar_eventos, pasta_temporaria, preparar_qt,
                       pular, resumir, secao, usar_fonte_windows)
import fixtures_docx as fx

QT = preparar_qt()

from editdoc.formato_docx import conferencia, ordem  # noqa: E402
from editdoc.formato_docx.escritor import escrever, escrever_com_avisos  # noqa: E402
from editdoc.formato_docx.leitor import LIMITES_PADRAO, ler  # noqa: E402
from editdoc.modelo.documento import (Celula, Documento, Linha, Paragrafo, Run, Tabela,  # noqa: E402
                                      normalizar, paragrafos_de, tabelas_de)

PASTA_AMOSTRAS = RAIZ / "tests" / "saida"
LIMITE_S = 2.0


def iguais(a: Documento, b: Documento) -> bool:
    return normalizar(a).paragrafos == normalizar(b).paragrafos


def textos(blocos) -> list[str]:
    return ["".join(r.texto for r in p.runs) for p in paragrafos_de(blocos)]


def docx_com_corpo(corpo_xml: str, estilos_extra: str = "") -> bytes:
    """Um .docx valido cujo corpo e' `corpo_xml`; `estilos_extra` entra no styles.xml."""
    base = fx.bytes_de(Documento(paragrafos=[Paragrafo(runs=[Run("x")])]))
    partes = fx.partes_de(base)
    xml = partes["word/document.xml"].decode("utf-8")
    inicio = xml.index("<w:body>") + len("<w:body>")
    fim = xml.index("<w:sectPr")
    partes["word/document.xml"] = (xml[:inicio] + corpo_xml + xml[fim:]).encode("utf-8")
    if estilos_extra:
        estilos = partes["word/styles.xml"].decode("utf-8")
        partes["word/styles.xml"] = estilos.replace("</w:styles>",
                                                    estilos_extra + "</w:styles>").encode("utf-8")
    return fx.montar(partes)


def _tc(texto: str, tcpr: str = "") -> str:
    return (f"<w:tc><w:tcPr>{tcpr}</w:tcPr><w:p><w:r><w:t>{texto}</w:t></w:r></w:p></w:tc>")


# -- modelo e .docx ---------------------------------------------------------------

def testar_ida_e_volta_docx() -> None:
    secao("modelo -> .docx -> modelo: mesclas, bordas, sombra, cabecalho, aninhada")
    casos = {
        "tabela mesclada": Documento(paragrafos=[fx.tabela_mesclada()]),
        "amostra completa do M4": fx.documento_com_tabela_e_imagem(),
        "tabela no comeco e texto depois": Documento(paragrafos=[
            fx.tabela_mesclada(), Paragrafo(runs=[Run("fim")])]),
        "duas tabelas seguidas": Documento(paragrafos=[fx.tabela_mesclada(), fx.tabela_mesclada()]),
    }
    for nome, doc in casos.items():
        dados, avisos_gravacao = escrever_com_avisos(doc)
        lido, avisos = ler(dados)
        checa(iguais(doc, lido), f"{nome}: modelo igual depois da volta")
        checa_igual(avisos, [], f"{nome}: o proprio arquivo nao gera aviso de leitura")
        checa_igual(avisos_gravacao, [], f"{nome}: a gravacao nao remove nada")
        checa_igual(ordem.validar(fx.partes_de(dados)), [], f"{nome}: o pacote passa no validador")
        checa(escrever(lido) == escrever(ler(escrever(lido))[0]),
              f"{nome}: a segunda ida produz os mesmos bytes (idempotente)")
    doc = casos["tabela mesclada"]
    lido, _ = ler(escrever(doc))
    tab = lido.paragrafos[0]
    checa_igual([c.colunas for c in tab.linhas[0].celulas], [2, 1], "gridSpan volta: 2 + 1 colunas")
    checa_igual([lin.celulas[0].mescla_vertical for lin in tab.linhas], [None, "inicio", "continua"],
                "vMerge volta: inicio e continua")
    checa_igual(tab.linhas[0].celulas[0].sombreamento_hex, "D9E2F3", "sombreamento volta")
    checa_igual(tab.bordas.sup.estilo, "dupla", "borda dupla de cima volta")
    checa_igual(tab.linhas[0].celulas[1].bordas.inf.cor_hex, "0070C0", "borda propria da celula volta")
    checa(tab.linhas[0].cabecalho and not tab.linhas[1].cabecalho, "linha de cabecalho volta")
    checa_igual(tab.linhas[1].altura_tw, 600, "altura da linha volta")
    checa_igual(tab.colunas_tw, [3000, 2500, 3500], "larguras da grade voltam")

    alterado = normalizar(lido)
    alterado.paragrafos[0].linhas[2].celulas[0].mescla_vertical = None
    checa(not iguais(doc, alterado), "contraprova: desfazer a mescla vertical e' acusado")
    alterado = normalizar(lido)
    alterado.paragrafos[0].linhas[0].celulas[0].colunas = 1
    checa(not iguais(doc, alterado), "contraprova: trocar o gridSpan e' acusado")


def testar_leitura_do_word() -> None:
    secao("leitura de tabela como o Word grava: estilo Table Grid, hMerge, linha curta")
    estilo = ('<w:style w:type="table" w:styleId="TabelaGrade"><w:name w:val="Table Grid"/>'
              '<w:tblPr><w:tblBorders><w:top w:val="single" w:sz="4" w:color="auto"/>'
              '<w:left w:val="single" w:sz="4"/><w:bottom w:val="single" w:sz="4"/>'
              '<w:right w:val="single" w:sz="4"/><w:insideH w:val="single" w:sz="4"/>'
              '<w:insideV w:val="single" w:sz="4"/></w:tblBorders></w:tblPr></w:style>')
    corpo = ('<w:tbl><w:tblPr><w:tblStyle w:val="TabelaGrade"/><w:tblW w:w="0" w:type="auto"/>'
             '<w:tblLook w:val="04A0"/></w:tblPr>'
             '<w:tblGrid><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/><w:gridCol w:w="2000"/></w:tblGrid>'
             '<w:tr>' + _tc("a", '<w:hMerge w:val="restart"/>') + _tc("", '<w:hMerge/>') + _tc("c") + '</w:tr>'
             '<w:tr>' + _tc("d") + _tc("e") + '</w:tr>'
             '</w:tbl><w:p/>')
    modelo, avisos = ler(docx_com_corpo(corpo, estilo))
    tab = modelo.paragrafos[0]
    checa(isinstance(tab, Tabela), "a tabela do arquivo vira Tabela, nao paragrafos")
    checa_igual(tab.bordas.vertical.estilo, "simples",
                "a borda do estilo Table Grid vale quando a tabela nao tem borda propria")
    checa_igual([c.colunas for c in tab.linhas[0].celulas], [2, 1],
                "hMerge (mescla horizontal antiga) vira gridSpan")
    checa_igual(sum(c.colunas for c in tab.linhas[1].celulas), 3,
                "linha mais curta que a grade e' completada com celula vazia")
    sem_estilo, _ = ler(docx_com_corpo(corpo.replace('<w:tblStyle w:val="TabelaGrade"/>', "")))
    checa_igual(sem_estilo.paragrafos[0].bordas.vertical.estilo, "nenhuma",
                "contraprova: sem estilo nem borda propria, a tabela nao tem borda")
    checa_igual(textos(modelo.paragrafos)[:4], ["a", "c", "d", "e"], "o texto das celulas volta")

    escondida = ('<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid>'
                 '<w:tr>' + _tc("topo", '<w:vMerge w:val="restart"/>') + '</w:tr>'
                 '<w:tr>' + _tc("escondido", "<w:vMerge/>") + '</w:tr></w:tbl><w:p/>')
    _m, avisos = ler(docx_com_corpo(escondida))
    checa(any(a.codigo == "CELULA_ESCONDIDA" for a in avisos),
          "texto em celula coberta por mescla vertical gera aviso (o Word nao o mostra)")
    quebra = ('<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="2000"/></w:tblGrid><w:tr><w:tc>'
              '<w:p><w:r><w:t>a</w:t><w:br w:type="page"/><w:t>b</w:t></w:r></w:p></w:tc></w:tr>'
              '</w:tbl><w:p/>')
    m_q, avisos = ler(docx_com_corpo(quebra))
    checa(any(a.codigo == "QUEBRA_EM_TABELA" for a in avisos)
          and textos(m_q.paragrafos)[0] == "a\nb",
          "quebra de pagina dentro da celula vira quebra de linha, com aviso")


def testar_oraculo() -> None:
    secao("python-docx como oraculo: mesmas celulas, mesmas mesclas, mesmo texto")
    import docx

    dados = fx.bytes_de(Documento(paragrafos=[fx.tabela_mesclada()]))
    d = docx.Document(io.BytesIO(dados))
    checa_igual(len(d.tables), 1, "python-docx ve uma tabela")
    t = d.tables[0]
    checa_igual((len(t.rows), len(t.columns)), (3, 3), "3 linhas e 3 colunas")
    checa(t.cell(0, 0)._tc is t.cell(0, 1)._tc, "a mescla horizontal e' uma celula so' para o Word")
    checa(t.cell(1, 0).text == "Mescla vertical" and t.cell(2, 0).text == "Mescla vertical",
          "a mescla vertical mostra o texto da primeira celula nas duas linhas")
    checa_igual(t.cell(2, 2).text, "C3", "texto de celula comum")
    checa_igual(t.style.name if t.style is not None else None, "Table Grid",
                "a tabela usa o estilo Table Grid (espacamento zero na celula)")
    larguras = [round(c.width.twips) if c.width is not None else None for c in t.rows[1].cells]
    checa_igual(larguras, [3000, 2500, 3500], "a largura de cada celula bate com a grade")


def testar_validador() -> None:
    secao("validador: celula sem paragrafo, linha torta e corpo terminado em tabela")
    dados = fx.bytes_de(Documento(paragrafos=[fx.tabela_mesclada()]))
    partes = fx.partes_de(dados)
    checa_igual(ordem.validar(partes), [], "o arquivo do EditDoc passa")
    xml = partes["word/document.xml"]
    sabotagens = {
        "tc sem paragrafo no fim": xml.replace(b"<w:p><w:r><w:t>C3</w:t></w:r></w:p></w:tc>",
                                               b"</w:tc>", 1),
        "linha com uma coluna a mais": xml.replace(b'<w:gridSpan w:val="2"/>',
                                                   b'<w:gridSpan w:val="3"/>', 1),
        "corpo termina em tabela": xml.replace(b"</w:tbl><w:p/>", b"</w:tbl>", 1),
        "tcPr fora de ordem": xml.replace(b'<w:tcW w:w="5500" w:type="dxa"/><w:gridSpan w:val="2"/>',
                                          b'<w:gridSpan w:val="2"/><w:tcW w:w="5500" w:type="dxa"/>', 1),
    }
    for nome, novo in sabotagens.items():
        checa(novo != xml, f"{nome}: a sabotagem mudou o XML")
        partes_sab = dict(partes)
        partes_sab["word/document.xml"] = novo
        checa(len(ordem.validar(partes_sab)) >= 1, f"{nome}: o validador acusa")


def testar_limites() -> None:
    secao("limites contra arquivo hostil: linhas, colunas, celulas e aninhamento")
    lim = LIMITES_PADRAO

    def tabela_xml(linhas: int, colunas: int, texto: str = "t") -> str:
        grade = "".join('<w:gridCol w:w="100"/>' for _ in range(colunas))
        linha = "<w:tr>" + "".join(f"<w:tc><w:p><w:r><w:t>{texto}</w:t></w:r></w:p></w:tc>"
                                   for _ in range(colunas)) + "</w:tr>"
        return f"<w:tbl><w:tblPr/><w:tblGrid>{grade}</w:tblGrid>{linha * linhas}</w:tbl><w:p/>"

    casos = {
        "linhas acima do limite": (tabela_xml(lim["tabela_linhas"] + 1, 1), lim["tabela_linhas"] + 1),
        "colunas acima do limite": (tabela_xml(1, lim["tabela_colunas"] + 1), lim["tabela_colunas"] + 1),
        "celulas acima do limite": (tabela_xml(lim["tabela_celulas"] // 4 + 1, 4),
                                    (lim["tabela_celulas"] // 4 + 1) * 4),
    }
    for nome, (xml, n_textos) in casos.items():
        inicio = time.perf_counter()
        modelo, avisos = ler(docx_com_corpo(xml))
        tempo = time.perf_counter() - inicio
        checa(not any(isinstance(b, Tabela) for b in modelo.paragrafos)
              and any(a.codigo == "TABELA" for a in avisos),
              f"{nome}: a tabela vira paragrafos, com aviso TABELA")
        checa(textos(modelo.paragrafos).count("t") == n_textos and tempo < LIMITE_S,
              f"{nome}: nenhum texto se perde, em {tempo:.2f} s")
    no_limite, avisos = ler(docx_com_corpo(tabela_xml(lim["tabela_celulas"] // 4, 4)))
    checa(isinstance(no_limite.paragrafos[0], Tabela) and not avisos,
          "contraprova: tabela exatamente no limite de celulas continua tabela, sem aviso")

    profunda = "<w:p/>"
    for _ in range(lim["tabela_profundidade"] + 2):
        profunda = ('<w:tbl><w:tblPr/><w:tblGrid><w:gridCol w:w="100"/></w:tblGrid><w:tr><w:tc>'
                    + profunda.replace("<w:p/>", "<w:p><w:r><w:t>fundo</w:t></w:r></w:p>", 1)
                    + "<w:p/></w:tc></w:tr></w:tbl><w:p/>")
    modelo, avisos = ler(docx_com_corpo(profunda))
    niveis = 0
    blocos = modelo.paragrafos
    while blocos and isinstance(blocos[0], Tabela):
        niveis += 1
        blocos = blocos[0].linhas[0].celulas[0].blocos
    checa_igual(niveis, lim["tabela_profundidade"], "aninhamento alem do limite vira paragrafos")
    checa(any(a.codigo == "TABELA" for a in avisos) and "fundo" in textos(modelo.paragrafos),
          "a tabela funda demais avisa e o texto dela fica")


# -- ponte Qt e tela ---------------------------------------------------------------

def testar_ponte_qt() -> None:
    secao("ponte Qt: mesclas viram spans, bloco estrutural nao volta, segunda volta igual")
    from PySide6.QtGui import QTextDocument

    from editdoc.interface.layout_paginado import LayoutPaginado
    from editdoc.modelo.pagina import Pagina
    from editdoc.ponte_qt import de_qt, para_qt

    casos = {
        "tabela no comeco": Documento(paragrafos=[fx.tabela_mesclada(), Paragrafo(runs=[Run("fim")])]),
        "amostra completa do M4": fx.documento_com_tabela_e_imagem(),
        "duas tabelas seguidas": Documento(paragrafos=[fx.tabela_mesclada(), fx.tabela_mesclada()]),
        "tabela so'": Documento(paragrafos=[fx.tabela_mesclada()]),
    }
    for nome, doc in casos.items():
        for qdoc in (QTextDocument(), LayoutPaginado(Pagina(), "Calibri", 11).doc):
            para_qt.preencher(qdoc, doc)
            volta = de_qt.extrair(qdoc)
            checa(iguais(doc, volta), f"{nome} ({type(qdoc).__name__}): igual depois da volta pelo Qt")
            para_qt.preencher(qdoc, volta)
            checa(iguais(volta, de_qt.extrair(qdoc)), f"{nome} ({type(qdoc).__name__}): segunda volta igual")
    qdoc = QTextDocument()
    para_qt.preencher(qdoc, Documento(paragrafos=[fx.tabela_mesclada()]))
    tabela = qdoc.rootFrame().childFrames()[0]
    checa_igual((tabela.rows(), tabela.columns()), (3, 3), "QTextTable 3 x 3")
    checa_igual((tabela.cellAt(0, 0).columnSpan(), tabela.cellAt(1, 0).rowSpan()), (2, 2),
                "gridSpan vira columnSpan e vMerge vira rowSpan")
    checa_igual(tabela.format().headerRowCount(), 1, "a linha de cabecalho repete no alto da pagina")
    primeiro = qdoc.begin()
    checa(bool(primeiro.blockFormat().property(para_qt.PROP_ESTRUTURAL)) and primeiro.text() == "",
          "o bloco que o Qt exige antes da tabela leva a marca estrutural")
    checa_igual(len([b for b in de_qt.extrair(qdoc).paragrafos if isinstance(b, Paragrafo)]), 0,
                "o bloco estrutural vazio nao volta como paragrafo")


def _painel(largura: int = 1000, altura: int = 800):
    from editdoc.interface.painel_documento import PainelDocumento
    painel = PainelDocumento("Calibri", 11)
    painel.resize(largura, altura)
    painel.show()
    drenar_eventos(2)
    painel.setFocus()
    return painel


def _contar_escuros(img, x0: int, y0: int, x1: int, y1: int) -> int:
    n = 0
    for y in range(max(0, y0), min(img.height(), y1)):
        for x in range(max(0, x0), min(img.width(), x1)):
            c = img.pixelColor(x, y)
            if c.red() < 90 and c.green() < 90 and c.blue() < 90:
                n += 1
    return n


def testar_tela() -> None:
    secao("tela: a tabela aparece com bordas; sem borda, sem tinta")
    from editdoc.modelo.documento import Borda, BordasTabela
    from editdoc.ponte_qt import para_qt

    painel = _painel()
    vazia = Tabela(colunas_tw=[3000, 3000], linhas=[
        Linha(celulas=[Celula(blocos=[Paragrafo()]), Celula(blocos=[Paragrafo()])])
        for _ in range(3)])
    para_qt.preencher(painel.document(), Documento(paragrafos=[vazia, Paragrafo()]))
    painel.clearFocus()  # o caret tambem e' tinta preta
    drenar_eventos(2)
    img = painel.viewport().grab().toImage()
    com_borda = _contar_escuros(img, 0, 0, img.width(), img.height() // 2)
    checa(com_borda > 300, f"as bordas da tabela tem tinta ({com_borda} px escuros)")
    sem = Tabela(colunas_tw=[3000, 3000], bordas=BordasTabela(*[Borda("nenhuma", 0)] * 6),
                 linhas=vazia.linhas)
    para_qt.preencher(painel.document(), Documento(paragrafos=[sem, Paragrafo()]))
    drenar_eventos(2)
    img = painel.viewport().grab().toImage()
    sem_borda = _contar_escuros(img, 0, 0, img.width(), img.height() // 2)
    checa(sem_borda < 5, f"contraprova: tabela sem borda nao tem tinta ({sem_borda} px)")

    # Cache de tiles: digitar numa celula ate' a linha crescer nao pode deixar tile velho.
    from PySide6.QtTest import QTest
    para_qt.preencher(painel.document(), Documento(paragrafos=[
        Paragrafo(runs=[Run("acima")]), vazia, Paragrafo(runs=[Run("abaixo")])]))
    drenar_eventos(2)
    painel.viewport().grab()  # enche o cache
    tabela = painel.document().rootFrame().childFrames()[0]
    painel.setTextCursor(tabela.cellAt(1, 0).firstCursorPosition())
    painel.setFocus()
    QTest.keyClicks(painel, "texto longo que quebra a linha da celula " * 3)
    painel.clearFocus()
    drenar_eventos(2)
    com_cache = painel.viewport().grab().toImage()
    painel._limpar_cache()
    novo = painel.viewport().grab().toImage()
    checa(com_cache == novo, "depois de digitar na celula, a tela com cache e' igual a um quadro novo")
    painel.document().setModified(False)
    painel.close()


def testar_paginas() -> None:
    secao("tabela entre paginas: nenhuma linha some, nenhum bloco fica cortado na divisa")
    from editdoc.ponte_qt import para_qt

    painel = _painel()
    longa = Tabela(colunas_tw=[4000, 4000], linhas=[
        Linha(cabecalho=(i == 0), celulas=[Celula(blocos=[Paragrafo(runs=[Run(f"linha {i}")])]),
                                           Celula(blocos=[Paragrafo(runs=[Run(f"valor {i}")])])])
        for i in range(120)])
    alta = Tabela(colunas_tw=[8000], linhas=[Linha(celulas=[Celula(blocos=[
        Paragrafo(runs=[Run(f"parte {i} de uma linha mais alta que a pagina")]) for i in range(150)])])])
    for nome, tabela in (("120 linhas", longa), ("uma linha mais alta que a pagina", alta)):
        para_qt.preencher(painel.document(), Documento(paragrafos=[tabela, Paragrafo()]))
        drenar_eventos(2)
        lay = painel.layout_pag
        dl = lay.doc.documentLayout()
        checa(painel.total_de_paginas() >= 3, f"{nome}: ocupa varias paginas ({painel.total_de_paginas()})")
        bloco = lay.doc.begin()
        cortados = 0
        tops = []
        while bloco.isValid():
            r = dl.blockBoundingRect(bloco)
            if bloco.text():
                tops.append(r.top())
                # Vale a linha de texto, nao o retangulo do bloco: o Qt empurra a linha
                # para a pagina seguinte com um deslocamento dentro do bloco (medido).
                lay_bloco = bloco.layout()
                for i in range(lay_bloco.lineCount()):
                    linha = lay_bloco.lineAt(i)
                    topo = r.top() + linha.y()
                    base = topo + linha.height()
                    if int(topo // lay.hb) != int((base - 1) // lay.hb):
                        cortados += 1
            bloco = bloco.next()
        checa_igual(cortados, 0, f"{nome}: nenhuma linha de texto fica cortada na divisa de pagina")
        checa(all(a <= b for a, b in zip(tops, tops[1:])), f"{nome}: os paragrafos seguem em ordem")
        checa(tops[-1] < lay.numero_paginas() * lay.hb, f"{nome}: o ultimo paragrafo esta' numa pagina")
        texto = painel.toPlainText()
        checa(("linha 119" in texto) if tabela is longa else ("parte 149" in texto),
              f"{nome}: o texto do fim da tabela continua no documento")
    painel.document().setModified(False)
    painel.close()


# -- edicao ---------------------------------------------------------------------

def _texto_celula(tabela, r: int, c: int) -> str:
    cel = tabela.cellAt(r, c)
    cur = cel.firstCursorPosition()
    cur.setPosition(cel.lastCursorPosition().position(), cur.MoveMode.KeepAnchor)
    return cur.selectedText()


def testar_edicao_teclado() -> None:
    secao("teclado: digitar na celula, Tab, Shift+Tab, Tab na ultima celula cria linha")
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest

    from editdoc.edicao import tabelas

    painel = _painel()
    QTest.keyClicks(painel, "antes")
    tabela = tabelas.inserir_tabela(painel, 2, 2)
    checa(tabela is not None and painel.textCursor().currentTable() is tabela,
          "inserir tabela poe o cursor na primeira celula")
    QTest.keyClicks(painel, "um")
    QTest.keyClick(painel, Qt.Key.Key_Tab)
    QTest.keyClicks(painel, "dois")
    checa_igual((_texto_celula(tabela, 0, 0), _texto_celula(tabela, 0, 1)), ("um", "dois"),
                "Tab leva para a proxima celula e o texto entra nela")
    QTest.keyClick(painel, Qt.Key.Key_Backtab, Qt.KeyboardModifier.ShiftModifier)
    checa(painel.textCursor().selectedText() == "um",
          "Shift+Tab volta para a celula anterior e seleciona o conteudo (como o Word)")
    QTest.keyClick(painel, Qt.Key.Key_Tab)
    QTest.keyClick(painel, Qt.Key.Key_Tab)
    QTest.keyClick(painel, Qt.Key.Key_Tab)
    checa_igual(tabela.rows(), 2, "Tab ate' a ultima celula nao cria linha")
    QTest.keyClick(painel, Qt.Key.Key_Tab)
    checa_igual(tabela.rows(), 3, "Tab na ultima celula cria uma linha no fim")
    cel = tabela.cellAt(painel.textCursor())
    checa_igual((cel.row(), cel.column()), (2, 0), "o cursor vai para a primeira celula da linha nova")
    painel.undo()
    checa_igual(tabela.rows(), 2, "um desfazer tira a linha criada pelo Tab")
    checa("\t" not in painel.toPlainText(), "Tab na tabela nao insere tabulacao")
    QTest.keyClick(painel, Qt.Key.Key_Tab, Qt.KeyboardModifier.ControlModifier)
    checa("\t" in painel.toPlainText(), "Ctrl+Tab insere tabulacao dentro da celula")
    paginas = painel.total_de_paginas()
    QTest.keyClick(painel, Qt.Key.Key_Return, Qt.KeyboardModifier.ControlModifier)
    drenar_eventos(1)
    checa_igual(painel.total_de_paginas(), paginas,
                "Ctrl+Enter dentro da celula nao cria pagina (vira Enter comum)")
    checa("antes" in painel.toPlainText().split("\n")[0], "o texto antes da tabela fica acima dela")
    painel.document().setModified(False)
    painel.close()


def testar_comandos() -> None:
    secao("comandos: linha, coluna, mescla, divisao, exclusao; um desfazer cada")
    from PySide6.QtGui import QTextCursor

    from editdoc.edicao import tabelas

    painel = _painel()
    painel.insertPlainText("acima")
    tabela = tabelas.inserir_tabela(painel, 2, 3)
    painel.undo()
    checa_igual(painel.document().rootFrame().childFrames(), [], "um desfazer tira a tabela inserida")
    checa_igual(painel.toPlainText(), "acima", "o desfazer devolve o texto como era")
    painel.moveCursor(painel.textCursor().MoveOperation.End)
    checa(tabelas.inserir_tabela(painel, 3000, 3) is None,
          "tabela acima do limite de celulas nao e' criada")
    tabela = tabelas.inserir_tabela(painel, 2, 3)
    for r in range(2):
        for c in range(3):
            cur = tabela.cellAt(r, c).firstCursorPosition()
            cur.insertText(f"{r}{c}")
    larguras = sum(r.rawValue() for r in tabela.format().columnWidthConstraints())
    corpo_px = painel.configuracao_pagina().corpo_largura_tw() / 15
    checa(abs(larguras - corpo_px) < 3, "a tabela nova ocupa a largura do corpo")

    def em(r: int, c: int) -> None:
        painel.setTextCursor(tabela.cellAt(r, c).firstCursorPosition())

    def um_desfazer(acao, descricao: str, medida) -> None:
        antes = medida()
        acao()
        depois = medida()
        painel.undo()
        volta = medida()
        painel.redo()
        refeito = medida()
        checa(depois != antes and volta == antes and refeito == depois,
              f"{descricao}: muda, um desfazer devolve e refazer repete ({antes} -> {depois})")

    em(0, 1)
    um_desfazer(lambda: tabelas.inserir_linha(painel, abaixo=True), "inserir linha abaixo",
                lambda: tabela.rows())
    em(0, 1)
    um_desfazer(lambda: tabelas.inserir_linha(painel, abaixo=False), "inserir linha acima",
                lambda: (tabela.rows(), _texto_celula(tabela, 0, 0)))
    em(0, 1)
    um_desfazer(lambda: tabelas.inserir_coluna(painel, direita=True), "inserir coluna a direita",
                lambda: tabela.columns())
    total = sum(r.rawValue() for r in tabela.format().columnWidthConstraints())
    checa(abs(total - larguras) < 1, "a coluna nova nao alarga a tabela (larguras reduzidas)")
    em(0, 0)
    um_desfazer(lambda: tabelas.excluir_colunas(painel), "excluir coluna", lambda: tabela.columns())
    em(1, 0)
    um_desfazer(lambda: tabelas.excluir_linhas(painel), "excluir linha", lambda: tabela.rows())

    cur = tabela.cellAt(0, 0).firstCursorPosition()
    cur.setPosition(tabela.cellAt(1, 1).lastCursorPosition().position(), QTextCursor.MoveMode.KeepAnchor)
    painel.setTextCursor(cur)
    checa(tabelas.pode_mesclar(painel.textCursor()), "selecao de 2 x 2 celulas pode mesclar")
    um_desfazer(lambda: tabelas.mesclar_celulas(painel), "mesclar celulas",
                lambda: (tabela.cellAt(0, 0).rowSpan(), tabela.cellAt(0, 0).columnSpan()))
    em(0, 0)
    checa(tabelas.pode_dividir(painel.textCursor()), "a celula mesclada pode ser dividida")
    um_desfazer(lambda: tabelas.dividir_celula(painel), "dividir celula",
                lambda: (tabela.cellAt(0, 0).rowSpan(), tabela.cellAt(0, 0).columnSpan()))

    em(0, 0)
    painel.moveCursor(QTextCursor.MoveOperation.End)
    painel.insertPlainText("abaixo")
    em(0, 0)
    quadros = lambda: len(painel.document().rootFrame().childFrames())  # noqa: E731
    um_desfazer(lambda: tabelas.excluir_tabela(painel), "excluir tabela", quadros)
    em = None  # a tabela pode ter sido recriada pelo refazer: nao reusar a referencia
    blocos = []
    bloco = painel.document().begin()
    while bloco.isValid():
        blocos.append(bloco.text())
        bloco = bloco.next()
    checa(blocos[0] == "acima" and "abaixo" in blocos and len(blocos) >= 2,
          f"excluir tabela nao junta o paragrafo de cima com o de baixo ({blocos})")
    painel.document().setModified(False)
    painel.close()


def testar_colar() -> None:
    secao("colar: tabela copiada e tabela de HTML de fora viram tabela do EditDoc")
    from PySide6.QtCore import QMimeData
    from PySide6.QtGui import QTextCursor
    from PySide6.QtWidgets import QApplication

    from editdoc.edicao import tabelas
    from editdoc.ponte_qt import de_qt, para_qt

    painel = _painel()
    para_qt.preencher(painel.document(), Documento(paragrafos=[
        Paragrafo(runs=[Run("antes")]), fx.tabela_mesclada(), Paragrafo(runs=[Run("depois")])]))
    tabela = painel.document().rootFrame().childFrames()[0]
    cur = QTextCursor(painel.document())
    cur.setPosition(tabela.firstPosition() - 1)
    cur.setPosition(tabela.lastPosition() + 1, QTextCursor.MoveMode.KeepAnchor)
    painel.setTextCursor(cur)
    painel.copy()
    painel.moveCursor(QTextCursor.MoveOperation.End)
    painel.paste()
    tabelas_qt = tabelas.tabelas_do_documento(painel.document())
    checa_igual(len(tabelas_qt), 2, "a tabela copiada e colada vira uma segunda tabela")
    checa(all(t.format().hasProperty(para_qt.PROP_TABELA) for t in tabelas_qt),
          "a tabela colada tem o formato do EditDoc (borda colapsada, largura fixa)")
    modelo = de_qt.extrair(painel.document())
    grades = [(len(t.linhas), len(t.colunas_tw)) for t in tabelas_de(modelo.paragrafos)]
    checa_igual(grades, [(3, 3), (3, 3)], "a copia tem a mesma grade")

    md = QMimeData()
    md.setHtml("<table border='1'><tr><td>x</td><td>y</td></tr><tr><td>z</td><td>w</td></tr></table>")
    QApplication.clipboard().setMimeData(md)
    painel.moveCursor(QTextCursor.MoveOperation.End)
    painel.paste()
    modelo = de_qt.extrair(painel.document())
    ultima = list(tabelas_de(modelo.paragrafos))[-1]
    checa_igual((len(ultima.linhas), len(ultima.colunas_tw)), (2, 2), "a tabela de HTML entra 2 x 2")
    corpo = painel.configuracao_pagina().corpo_largura_tw()
    checa(abs(sum(ultima.colunas_tw) - corpo) <= 2,
          "a tabela de HTML, sem largura, ganha a largura do corpo")
    checa(all(t.format().hasProperty(para_qt.PROP_TABELA)
              for t in tabelas.tabelas_do_documento(painel.document())),
          "a tabela de HTML tambem vira tabela do EditDoc")
    painel.undo()
    checa_igual(len(tabelas.tabelas_do_documento(painel.document())), 2,
                "um desfazer tira a tabela colada (com o ajuste junto)")
    painel.document().setModified(False)
    painel.close()


def testar_faixa_e_janela(tmp: pathlib.Path) -> None:
    secao("janela: grade da aba Inserir, dialogo, comandos habilitados so' na tabela, salvar")
    from PySide6.QtCore import QEvent, QPointF, Qt
    from PySide6.QtGui import QMouseEvent
    from PySide6.QtWidgets import QApplication

    from editdoc import configuracao
    from editdoc.interface import dialogos_modais
    from editdoc.interface.janela_principal import JanelaPrincipal

    usar_fonte_windows(9)
    janela = JanelaPrincipal(configuracao.padrao())
    janela.show()
    drenar_eventos()
    c = janela.comandos
    checa(not c.acao("inserir_linha_abaixo").isEnabled(), "fora da tabela, inserir linha fica desabilitado")
    janela.faixa.abrir_aba("Inserir")
    seletor = janela.faixa.seletor_grade
    ponto = seletor.centro_da_celula(3, 4)
    for tipo in (QEvent.Type.MouseMove, QEvent.Type.MouseButtonRelease):
        ev = QMouseEvent(tipo, ponto, seletor.mapToGlobal(ponto), Qt.MouseButton.LeftButton,
                         Qt.MouseButton.LeftButton if tipo == QEvent.Type.MouseMove else Qt.MouseButton.NoButton,
                         Qt.KeyboardModifier.NoModifier)
        QApplication.sendEvent(seletor, ev)
    drenar_eventos()
    tabela = janela.painel.textCursor().currentTable()
    checa(tabela is not None and (tabela.rows(), tabela.columns()) == (3, 4),
          "clicar na grade em 3 x 4 insere uma tabela de 3 linhas e 4 colunas")
    checa(c.acao("inserir_linha_abaixo").isEnabled() and c.acao("excluir_tabela").isEnabled(),
          "com o cursor na tabela, os comandos de tabela habilitam")
    checa(not c.acao("mesclar_celulas").isEnabled(), "sem selecao de celulas, mesclar fica desabilitado")
    c.acao("inserir_linha_abaixo").trigger()
    checa_igual(tabela.rows(), 4, "o botao Inserir linha abaixo funciona pela acao")

    original = dialogos_modais.pedir_tamanho_tabela
    try:
        dialogos_modais.pedir_tamanho_tabela = lambda *_a: (2, 5)
        janela.painel.moveCursor(janela.painel.textCursor().MoveOperation.End)
        drenar_eventos()
        c.acao("inserir_tabela").trigger()
    finally:
        dialogos_modais.pedir_tamanho_tabela = original
    segunda = janela.painel.textCursor().currentTable()
    checa(segunda is not None and segunda is not tabela and
          (segunda.rows(), segunda.columns()) == (2, 5),
          "o dialogo Inserir tabela cria a tabela pedida (2 x 5)")

    destino = tmp / "tabela.docx"
    with_dialogo = dialogos_modais.pedir_salvar_arquivo
    try:
        dialogos_modais.pedir_salvar_arquivo = lambda *_a: str(destino)
        checa(janela.salvar_como(), "salvar o documento com duas tabelas")
    finally:
        dialogos_modais.pedir_salvar_arquivo = with_dialogo
    modelo, avisos = ler(destino.read_bytes())
    checa_igual([(len(t.linhas), len(t.colunas_tw)) for t in tabelas_de(modelo.paragrafos)],
                [(4, 4), (2, 5)], "o arquivo gravado tem as duas tabelas com as mesmas grades")
    checa_igual(avisos, [], "o arquivo gravado nao gera aviso ao abrir")
    checa_igual(ordem.validar(fx.partes_de(destino.read_bytes())), [], "o arquivo gravado passa no validador")
    janela.painel.document().setModified(False)
    janela.close()


def testar_salvar_e_amostra(tmp: pathlib.Path) -> None:
    secao("abrir, editar celula mesclada, salvar e reabrir; amostra para o Word")
    from editdoc import configuracao
    from editdoc.interface.gerente import GerenteDeJanelas

    caminho = tmp / "m4.docx"
    caminho.write_bytes(fx.bytes_de(fx.documento_com_tabela_e_imagem()))
    gerente = GerenteDeJanelas(configuracao.padrao())
    janela = gerente.abrir(str(caminho))
    checa(janela is not None and janela.barra_aviso.isHidden(),
          "o arquivo com tabela mesclada e imagem abre sem aviso")
    tabela = janela.painel.document().rootFrame().childFrames()[0]
    cur = tabela.cellAt(1, 0).firstCursorPosition()
    janela.painel.setTextCursor(cur)
    janela.painel.insertPlainText("editado ")
    checa(janela.salvar(), "salvar grava (a conferencia aceita a tabela)")
    lido, avisos = ler(caminho.read_bytes())
    tab = next(tabelas_de(lido.paragrafos))
    checa(textos(tab.linhas[1].celulas[0].blocos)[0].startswith("editado "),
          "o texto digitado na celula mesclada esta' no arquivo")
    checa_igual([lin.celulas[0].mescla_vertical for lin in tab.linhas], [None, "inicio", "continua"],
                "a mescla vertical sobrevive a abrir, editar e salvar")
    checa_igual(avisos, [], "reabrir nao gera aviso")
    esperado = normalizar(fx.documento_com_tabela_e_imagem())
    esperado_tab = next(tabelas_de(esperado.paragrafos))
    esperado_tab.linhas[1].celulas[0].blocos[0].runs.insert(0, Run("editado "))
    checa(iguais(esperado, lido), "o resto do documento volta igual ao original")
    try:
        conferencia.conferir(caminho.read_bytes(), Documento(paragrafos=[fx.tabela_mesclada()]))
        acusou = False
    except conferencia.FalhaDeConferencia:
        acusou = True
    checa(acusou, "contraprova: a conferencia acusa um arquivo com estrutura diferente")
    janela.painel.document().setModified(False)
    janela.close()
    drenar_eventos()

    PASTA_AMOSTRAS.mkdir(parents=True, exist_ok=True)
    amostra = PASTA_AMOSTRAS / "amostra_m4_tabelas_imagens.docx"
    amostra.write_bytes(fx.bytes_de(fx.documento_com_tabela_e_imagem()))
    checa(amostra.is_file() and ordem.validar(fx.partes_de(amostra.read_bytes())) == [],
          f"amostra para abrir no Word: {amostra.relative_to(RAIZ)}")


def main() -> int:
    testar_ida_e_volta_docx()
    testar_leitura_do_word()
    testar_oraculo()
    testar_validador()
    testar_limites()
    if not QT:
        return pular("PySide6 nao esta instalado neste venv")
    if not usar_fonte_windows(9):
        return pular("Segoe UI nao esta instalada")
    testar_ponte_qt()
    testar_tela()
    testar_paginas()
    testar_edicao_teclado()
    testar_comandos()
    testar_colar()
    with pasta_temporaria("editdoc-teste-tabelas-") as tmp:
        testar_faixa_e_janela(tmp)
        testar_salvar_e_amostra(tmp)
    return resumir()


if __name__ == "__main__":
    raise SystemExit(main())
