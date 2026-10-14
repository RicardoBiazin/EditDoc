"""Campos PAGE e NUMPAGES (M5): numero, leitura, gravacao, ponte Qt e valor na tela.

- funcoes puras: formato do numero (decimal, romano, letras), opcao \\* da instrucao,
  numero inicial, variante da pagina (primeira, par, impar) e corpo empurrado pelo
  cabecalho;
- leitor: campo complexo (fldChar) e simples (fldSimple) viram campo; desde a 0.1.3 DATE,
  o campo que atravessa paragrafos e o IF com PAGE dentro sao preservados (marcas de inicio
  e fim, sem aviso); campo sem fim e campo aninhado fundo demais avisam (CAMPO);
- escritor: o campo sai como campo complexo, passa no validador, e o validador acusa
  campo sem fim; ida e volta e idempotencia; a conferencia acusa campo que virou texto;
- python-docx como oraculo: o Word ve o resultado e a instrucao;
- ponte Qt: o campo vira objeto e volta campo; digitar depois dele nao cria campo;
  inserir campo e' um passo de desfazer;
- estilo de caractere do campo (0.1.11): passa pela ponte Qt e o arquivo grava (o resultado do
  campo usava o numero de propriedade do estilo: o rodape "Pagina 1 de 3" em "Numero de pagina",
  num documento com estilo de id "1", nao salvava); nenhum numero de propriedade repetido;
- tela: o mesmo campo mostra 1, 2, 3... em cada pagina, e o total muda quando a
  paginacao muda.
"""

from __future__ import annotations

import io
import re

from ajudantes import (checa, checa_igual, drenar_eventos, preparar_qt, pular, resumir, secao,
                       usar_fonte_windows)
import fixtures_docx as fx

QT = preparar_qt()

from editdoc.formato_docx import conferencia, ordem  # noqa: E402
from editdoc.formato_docx.escritor import escrever, escrever_com_avisos  # noqa: E402
from editdoc.formato_docx.leitor import ler  # noqa: E402
from editdoc.modelo.campos import (formatar_numero, formato_da_instrucao, tipo_da_pagina,  # noqa: E402
                                   tipo_do_campo, valor_do_campo)
from editdoc.modelo.documento import Documento, Paragrafo, Run, normalizar  # noqa: E402
from editdoc.modelo.pagina import Pagina, pagina_efetiva  # noqa: E402


def codigos(avisos) -> dict[str, int]:
    return {a.codigo: a.contagem for a in avisos}


def runs_do_primeiro(modelo: Documento) -> list[Run]:
    return modelo.paragrafos[0].runs


# -- funcoes puras ------------------------------------------------------------------

def testar_numeros() -> None:
    secao("numero da pagina: formatos, opcao da instrucao, numero inicial e variante")
    casos = [(4, "decimal", "4"), (4, "lowerRoman", "iv"), (1994, "upperRoman", "MCMXCIV"),
             (9, "lowerRoman", "ix"), (1, "lowerLetter", "a"), (26, "upperLetter", "Z"),
             (27, "lowerLetter", "aa"), (53, "lowerLetter", "aaa"), (0, "lowerRoman", "0"),
             (4000, "upperRoman", "4000")]
    for n, formato, esperado in casos:
        checa_igual(formatar_numero(n, formato), esperado, f"{n} em {formato} = {esperado}")
    checa_igual(tipo_do_campo("page \\* MERGEFORMAT"), "PAGE", "PAGE em minusculas, com opcao")
    checa_igual(tipo_do_campo("NUMPAGES"), "NUMPAGES", "NUMPAGES")
    checa(tipo_do_campo("DATE \\@ dd/MM/yyyy") is None, "DATE nao e' campo de pagina")
    checa(tipo_do_campo("PAGE " + "x" * 300) is None, "instrucao longa demais nao e' aceita")
    checa(tipo_do_campo("PAGEREF _Toc1") is None, "PAGEREF nao e' PAGE (contraprova do prefixo)")
    checa_igual(formato_da_instrucao("PAGE \\* ROMAN \\* MERGEFORMAT"), "upperRoman",
                "opcao \\* ROMAN = romanos maiusculos")
    pagina = Pagina(formato_numero="lowerRoman", numero_inicial=5)
    checa_igual(valor_do_campo("PAGE", 0, 10, pagina), "v", "PAGE na 1a pagina com inicio 5, romano")
    checa_igual(valor_do_campo("PAGE", 2, 10, pagina), "vii", "PAGE na 3a pagina = vii")
    checa_igual(valor_do_campo("PAGE \\* Arabic", 2, 10, pagina), "7", "a opcao da instrucao vence")
    checa_igual(valor_do_campo("NUMPAGES", 2, 10, pagina), "10",
                "NUMPAGES conta paginas, sem numero inicial nem formato da secao")
    p = Pagina(primeira_diferente=True, pares_impares=True)
    checa_igual([tipo_da_pagina(i, p) for i in range(5)],
                ["primeira", "par", "padrao", "par", "padrao"],
                "primeira, par, impar, par, impar")
    p2 = Pagina(pares_impares=True, numero_inicial=2)
    checa_igual(tipo_da_pagina(0, p2), "par", "par ou impar segue o numero mostrado (inicio 2)")
    checa_igual([tipo_da_pagina(i, Pagina()) for i in range(3)], ["padrao"] * 3,
                "sem opcoes, toda pagina usa o padrao")


def testar_corpo_empurrado() -> None:
    secao("corpo empurrado: cabecalho alto desce o corpo; rodape alto o encurta")
    p = Pagina()
    checa(pagina_efetiva(p, 400, 300) is p, "cabecalho que cabe na margem nao muda nada")
    e = pagina_efetiva(p, 1200, 0)
    checa_igual(e.margem_sup_tw, 720 + 1200, "cabecalho de 1200 tw a 720 tw do topo: corpo em 1920")
    e = pagina_efetiva(p, 0, 1000)
    checa_igual(e.margem_inf_tw, 720 + 1000, "rodape de 1000 tw: margem de baixo em 1720")
    e = pagina_efetiva(p, 20000, 20000)
    checa(e.corpo_altura_tw() >= 1134, f"cabecalho absurdo: o corpo fica com o minimo ({e.corpo_altura_tw()})")
    checa_igual(p.margem_sup_tw, 1440, "a pagina do arquivo nao muda")


# -- leitor ---------------------------------------------------------------------------

def testar_leitura() -> None:
    secao("leitura: fldChar e fldSimple viram campo; os outros sao preservados (0.1.3)")
    corpo = ("<w:p><w:r><w:t xml:space=\"preserve\">Página </w:t></w:r>"
             + fx.campo_complexo("PAGE   \\* MERGEFORMAT", "3", "<w:rPr><w:b/></w:rPr>")
             + "<w:r><w:t xml:space=\"preserve\"> de </w:t></w:r>"
             + '<w:fldSimple w:instr=" NUMPAGES "><w:r><w:t>9</w:t></w:r></w:fldSimple></w:p>')
    modelo, avisos = ler(fx.docx_com_xml(corpo))
    runs = runs_do_primeiro(modelo)
    checa_igual([(r.texto, r.campo) for r in runs],
                [("Página ", None), ("3", "PAGE \\* MERGEFORMAT"), (" de ", None), ("9", "NUMPAGES")],
                "campo complexo e simples viram runs de campo, com o ultimo resultado")
    checa(runs[1].negrito is True, "o campo fica com a formatacao do resultado (negrito)")
    checa_igual(avisos, [], "campo de pagina nao gera aviso")

    # 0.1.3: DATE, REF, TOC... deixaram de virar texto fixo com o aviso CAMPO. As tres
    # checagens abaixo eram "DATE fica so' como texto", "DATE gera CAMPO" e "CAMPO pede
    # confirmacao"; passam a provar o campo preservado (marcas no fluxo, o resultado como
    # texto) e que o CAMPO, agora so' para o campo que nao pode ser preservado, continua
    # pedindo confirmacao. A ida e volta inteira esta' em teste_campos_gerais.py.
    from editdoc.formato_docx.avisos import construir
    from editdoc.modelo.documento import MarcaCampo
    corpo = "<w:p>" + fx.campo_complexo("DATE \\@ \"dd/MM/yyyy\"", "12/10/2026") + "</w:p>"
    modelo, avisos = ler(fx.docx_com_xml(corpo))
    runs = runs_do_primeiro(modelo)
    checa_igual([(r.texto, r.campo, type(r.marca).__name__ if r.marca else None) for r in runs],
                [("", None, "MarcaCampo"), ("12/10/2026", None, None), ("", None, "MarcaCampo")],
                "DATE vira campo preservado: inicio, o ultimo resultado como texto e fim")
    checa(isinstance(runs[0].marca, MarcaCampo) and runs[0].marca.campo.tipo() == "DATE"
          and runs[-1].marca.tipo == "fim", "a marca de inicio leva o codigo DATE; a outra e' o fim")
    checa_igual(codigos(avisos), {}, "DATE nao gera mais o aviso CAMPO (o campo volta ao arquivo)")
    checa(construir("CAMPO", 1).perde_ao_salvar,
          "o aviso CAMPO (campo que nao pode ser preservado) continua pedindo confirmacao")

    corpo = ("<w:p><w:r><w:fldChar w:fldCharType=\"begin\"/></w:r><w:r><w:instrText> PAGE "
             "</w:instrText></w:r><w:r><w:fldChar w:fldCharType=\"separate\"/></w:r>"
             "<w:r><w:t>um</w:t></w:r></w:p><w:p><w:r><w:t>dois</w:t></w:r>"
             "<w:r><w:fldChar w:fldCharType=\"end\"/></w:r></w:p>")
    modelo, avisos = ler(fx.docx_com_xml(corpo))
    textos = ["".join(r.texto for r in p.runs) for p in modelo.paragrafos]
    # 0.1.3: o campo que atravessa paragrafos e' preservado (inicio no 1o, fim no 2o).
    checa(textos[:2] == ["um", "dois"] and not any(r.campo for p in modelo.paragrafos for r in p.runs)
          and modelo.paragrafos[0].runs[0].marca.tipo == "inicio"
          and modelo.paragrafos[1].runs[-1].marca.tipo == "fim",
          "campo que atravessa paragrafos: o texto fica, e o campo vira inicio e fim em paragrafos diferentes")
    checa_igual(codigos(avisos), {}, "e nao gera mais o aviso CAMPO")

    corpo = "<w:p><w:r><w:fldChar w:fldCharType=\"begin\"/></w:r><w:r><w:t>resto</w:t></w:r></w:p>"
    modelo, avisos = ler(fx.docx_com_xml(corpo))
    checa_igual(codigos(avisos), {"CAMPO": 1}, "campo sem fim (arquivo cortado) gera aviso")
    checa_igual(runs_do_primeiro(modelo)[0].texto, "resto", "e o texto continua")

    profundo = "<w:p>" + "<w:r><w:fldChar w:fldCharType=\"begin\"/></w:r>" * 20 + \
               "<w:r><w:t>dentro</w:t></w:r>" + "<w:r><w:fldChar w:fldCharType=\"end\"/></w:r>" * 20 + "</w:p>"
    modelo, avisos = ler(fx.docx_com_xml(profundo))
    # 0.1.3: antes o texto "dentro" ficava solto. Ele esta' no CODIGO dos campos (nao ha'
    # separate), que o Word nao mostra; os 8 niveis do limite sao preservados e os outros 12
    # avisam. A checagem passa a ser: nao trava, avisa, e o arquivo gravado e' valido.
    dados, _av = escrever_com_avisos(modelo)
    checa(codigos(avisos).get("CAMPO", 0) >= 1 and ordem.validar(fx.partes_de(dados)) == [],
          "20 campos aninhados (limite 8): nao trava, ha' aviso e o arquivo gravado e' valido")

    aninhado = ("<w:p>" + "<w:r><w:fldChar w:fldCharType=\"begin\"/></w:r>"
                "<w:r><w:instrText> IF 1 = 1 </w:instrText></w:r>"
                + fx.campo_complexo("PAGE", "2")
                + "<w:r><w:fldChar w:fldCharType=\"separate\"/></w:r><w:r><w:t>sim</w:t></w:r>"
                "<w:r><w:fldChar w:fldCharType=\"end\"/></w:r></w:p>")
    modelo, avisos = ler(fx.docx_com_xml(aninhado))
    # 0.1.3: era "o IF nao e' entendido e avisa"; agora o IF volta com o PAGE dentro do codigo.
    codigo = runs_do_primeiro(modelo)[0].marca.campo.codigo
    checa("CAMPO" not in codigos(avisos) and codigo[0] == " IF 1 = 1 "
          and codigo[1].instrucao().strip() == "PAGE" and codigo[1].resultado == "2",
          "IF com PAGE dentro: o IF e' preservado, com o PAGE (e o resultado dele) no codigo")
    checa_igual("".join(r.texto for r in runs_do_primeiro(modelo) if not r.campo), "sim",
                "o resultado do IF fica como texto")


# -- escritor -------------------------------------------------------------------------

def _doc_com_campos() -> Documento:
    return Documento(pagina=Pagina(formato_numero="upperRoman", numero_inicial=3), paragrafos=[
        Paragrafo(runs=[Run("Página "), Run("1", campo="PAGE", negrito=True, tamanho_meio_pt=28),
                        Run(" de "), Run("7", campo="NUMPAGES \\* roman")], alinhamento="centro"),
        Paragrafo(runs=[Run("2", campo="PAGE"), Run("1", campo="PAGE")]),
        Paragrafo(runs=[Run("", campo="PAGE")]),
    ])


def testar_escrita() -> None:
    secao("escrita: campo complexo, validador, ida e volta, idempotencia e conferencia")
    doc = _doc_com_campos()
    dados, avisos = escrever_com_avisos(doc)
    checa_igual(avisos, [], "gravar campo nao gera aviso")
    xml = fx.partes_de(dados)["word/document.xml"].decode("utf-8")
    checa(xml.count('w:fldCharType="begin"') == 5 and xml.count('w:fldCharType="end"') == 5,
          "cada campo sai como begin ... end")
    checa('<w:instrText xml:space="preserve"> PAGE </w:instrText>' in xml,
          "a instrucao sai com xml:space=preserve")
    checa('<w:pgNumType w:fmt="upperRoman" w:start="3"/>' in xml, "formato e numero inicial no sectPr")
    checa_igual(ordem.validar(fx.partes_de(dados)), [], "o pacote passa no validador")
    lido, avisos_l = ler(dados)
    checa_igual(avisos_l, [], "o proprio arquivo abre sem aviso")
    checa(normalizar(lido).paragrafos == normalizar(doc).paragrafos, "o modelo volta igual (campos e resultados)")
    checa_igual((lido.pagina.formato_numero, lido.pagina.numero_inicial), ("upperRoman", 3),
                "formato e numero inicial voltam")
    checa(escrever(lido) == escrever(ler(escrever(lido))[0]), "a segunda ida produz os mesmos bytes")
    conferencia.conferir(dados, doc)
    checa(True, "a conferencia do arquivo gravado aceita os campos")

    # Contraprova da conferencia: o campo que vira texto e' acusado.
    sabotado = escrever(Documento(pagina=doc.pagina, paragrafos=[
        Paragrafo(runs=[Run("Página "), Run("1", negrito=True, tamanho_meio_pt=28), Run(" de "),
                        Run("7", campo="NUMPAGES \\* roman")], alinhamento="centro"),
        *doc.paragrafos[1:]]))
    try:
        conferencia.conferir(sabotado, doc)
        checa(False, "contraprova: o campo que virou texto passaria pela conferencia")
    except conferencia.FalhaDeConferencia as erro:
        checa("campo" in str(erro), f"contraprova: campo que virou texto e' acusado ({erro})")

    # Contraprova do validador: tirar o end deixa o campo aberto.
    partes = fx.partes_de(dados)
    partes["word/document.xml"] = xml.replace('<w:fldChar w:fldCharType="end"/>', "", 1).encode("utf-8")
    checa(any("fldChar" in p for p in ordem.validar(partes)), "contraprova: campo sem end e' acusado")
    partes["word/document.xml"] = xml.replace(
        '<w:instrText xml:space="preserve"> PAGE </w:instrText>', "<w:instrText> PAGE </w:instrText>", 1
    ).encode("utf-8")
    checa(any("instrText" in p for p in ordem.validar(partes)),
          "contraprova: instrucao com espaco sem xml:space e' acusada")


def testar_oraculo() -> None:
    secao("python-docx como oraculo: o Word ve o resultado e a instrucao")
    import docx
    dados = escrever(_doc_com_campos())
    d = docx.Document(io.BytesIO(dados))
    w = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    paragrafo = d.paragraphs[0]
    instrucoes = [e.text for e in paragrafo._p.iter(w + "instrText")]
    checa_igual(instrucoes, [" PAGE ", " NUMPAGES \\* roman "], "as instrucoes estao no paragrafo")
    checa_igual(paragrafo.text, "Página 1 de 7", "o texto visivel e' o ultimo resultado")
    checa(d.sections[0]._sectPr.find(w + "pgNumType") is not None, "pgNumType na secao")


# -- Qt ---------------------------------------------------------------------------------

def testar_ponte_qt() -> None:
    secao("ponte Qt: campo vira objeto e volta campo; digitar depois nao cria campo")
    from PySide6.QtGui import QTextDocument
    from PySide6.QtTest import QTest

    from editdoc.edicao import cabecalhos as ed_cabecalhos
    from editdoc.interface.painel_documento import PainelDocumento
    from editdoc.ponte_qt import campos, de_qt, para_qt

    doc = _doc_com_campos()
    qt = QTextDocument()
    para_qt.preencher(qt, doc)
    checa_igual(len(campos.posicoes_de_campos(qt)), 5, "cinco campos no QTextDocument")
    volta = de_qt.extrair(qt)
    volta.pagina = doc.pagina
    checa(normalizar(volta).paragrafos == normalizar(doc).paragrafos, "Qt -> modelo: os campos voltam iguais")

    painel = PainelDocumento("Calibri", 11)
    painel.resize(900, 700)
    painel.show()
    drenar_eventos()
    QTest.keyClicks(painel, "pag ")
    ed_cabecalhos.inserir_campo_no_cursor(painel, "PAGE")
    QTest.keyClicks(painel, " fim")
    modelo = de_qt.extrair(painel.document())
    checa_igual([(r.texto, r.campo) for r in modelo.paragrafos[0].runs],
                [("pag ", None), ("1", "PAGE"), (" fim", None)],
                "campo no cursor; o texto digitado depois nao vira campo")
    painel.setPlainText("")
    QTest.keyClicks(painel, "a")
    ed_cabecalhos.inserir_campo_no_cursor(painel, "NUMPAGES")
    painel.undo()
    checa_igual(painel.toPlainText(), "a", "inserir campo e' um passo de desfazer")
    painel.document().setModified(False)
    painel.close()


#: Estilos como o do documento real que nao salvava (0.1.11): ids numericos ("1" a "3", que o
#: Word gera para estilos de nome numerico) e o "Numero de pagina" de caractere.
#: Sem a lista e a revisao de ESTILOS_DO_WORD (pedem numbering.xml e geram aviso).
ESTILOS_COM_IDS_NUMERICOS = re.sub(r'<w:style [^>]*w:styleId="(?:Commarcadores|ComRevisao)".*?</w:style>',
                                   "", fx.ESTILOS_DO_WORD).replace("</w:styles>", "".join(
    f'<w:style w:type="paragraph" w:customStyle="1" w:styleId="{n}"><w:name w:val="{n}"/>'
    f'<w:basedOn w:val="Normal"/></w:style>' for n in ("1", "2", "3")) + (
    '<w:style w:type="character" w:styleId="Nmerodepgina"><w:name w:val="page number"/>'
    '<w:basedOn w:val="Fontepargpadro"/></w:style>'
    '<w:style w:type="paragraph" w:styleId="Rodap"><w:name w:val="footer"/>'
    '<w:basedOn w:val="Normal"/></w:style></w:styles>'))


def docx_rodape_com_estilo_no_campo() -> bytes:
    """O caso do arquivo que nao salvava, gerado em codigo: rodape com tabela, e na celula
    "Pagina: 0" + PAGE ("1") + " de 0" + NUMPAGES ("3"), tudo em "Numero de pagina"; e um
    paragrafo do corpo com o campo sozinho no estilo (ai' a perda nao se esconde atras de
    outro run com o mesmo estilo)."""
    rpr = '<w:rPr><w:rStyle w:val="Nmerodepgina"/><w:sz w:val="16"/></w:rPr>'
    celula = ('<w:p><w:pPr><w:pStyle w:val="Rodap"/></w:pPr>'
              '<w:r><w:rPr><w:sz w:val="16"/></w:rPr><w:t>Página: 0</w:t></w:r>'
              + fx.campo_complexo("PAGE", "1", rpr)
              + f'<w:r>{rpr}<w:t xml:space="preserve"> de 0</w:t></w:r>'
              + fx.campo_complexo("NUMPAGES", "3", rpr) + "</w:p>")
    tabela = ('<w:tbl><w:tblPr><w:tblW w:w="5000" w:type="dxa"/></w:tblPr><w:tblGrid>'
              '<w:gridCol w:w="2500"/><w:gridCol w:w="2500"/></w:tblGrid><w:tr>'
              '<w:tc><w:tcPr><w:tcW w:w="2500" w:type="dxa"/></w:tcPr>'
              '<w:p><w:pPr><w:pStyle w:val="Rodap"/></w:pPr><w:r><w:t>POP-001</w:t></w:r></w:p></w:tc>'
              '<w:tc><w:tcPr><w:tcW w:w="2500" w:type="dxa"/></w:tcPr>' + celula + "</w:tc>"
              "</w:tr></w:tbl><w:p><w:pPr><w:pStyle w:val=\"Rodap\"/></w:pPr></w:p>")
    rodape = fx.parte_faixa(tabela, "ftr")
    corpo = ('<w:p><w:r><w:t>Corpo.</w:t></w:r></w:p>'
             '<w:p>' + fx.campo_complexo("PAGE", "2", rpr) + "</w:p>")
    return fx.docx_com_xml(
        corpo, '<w:footerReference w:type="default" r:id="rId71"/>',
        partes_extra={"word/footer7.xml": rodape,
                      "word/styles.xml": ESTILOS_COM_IDS_NUMERICOS.encode("utf-8")},
        relacoes_extra=[("rId71", fx.RT_FOOTER, "footer7.xml")],
        tipos_extra=[("/word/footer7.xml", fx.CT_FOOTER)])


def _estilos_dos_campos(modelo: Documento) -> list:
    from editdoc.modelo.documento import Tabela
    saida = []

    def varrer(blocos):
        for b in blocos:
            if isinstance(b, Tabela):
                for linha in b.linhas:
                    for celula in linha.celulas:
                        varrer(celula.blocos)
            else:
                saida.extend((r.campo, r.estilo) for r in b.runs if r.campo)
    varrer(modelo.paragrafos)
    for tipo in sorted(modelo.rodapes):
        varrer(modelo.rodapes[tipo])
    return saida


def testar_estilo_do_campo() -> None:
    secao("estilo de caractere do campo (0.1.11): passa pela ponte Qt e o arquivo grava")
    import pathlib
    import tempfile

    from editdoc import configuracao
    from editdoc.interface import dialogos_modais
    from editdoc.interface.janela_principal import JanelaPrincipal
    from editdoc.ponte_qt import campos
    from editdoc.ponte_qt import estilos as pq_estilos

    dados = docx_rodape_com_estilo_no_campo()
    modelo, avisos = ler(dados)
    checa_igual(avisos, [], "a fixture abre sem aviso")
    esperado = [("PAGE", "Nmerodepgina"), ("PAGE", "Nmerodepgina"), ("NUMPAGES", "Nmerodepgina")]
    checa_igual(_estilos_dos_campos(modelo), esperado, "o leitor ve o estilo de cada campo")
    checa(campos.PROP_CAMPO_RESULTADO != pq_estilos.PROP_ESTILO_CAR,
          "o resultado do campo nao usa o numero do estilo de caractere")

    falhas: list[str] = []
    originais = (dialogos_modais.avisar, dialogos_modais.pedir_falha_na_conferencia)
    dialogos_modais.avisar = lambda _p, _t, texto: falhas.append(texto)
    dialogos_modais.pedir_falha_na_conferencia = lambda _p, nome, _l: falhas.append(nome) or "fechar"

    def salvar_pela_janela(pasta: pathlib.Path) -> tuple[bool, list]:
        janela = JanelaPrincipal(configuracao.padrao())
        janela.carregar(ler(dados)[0], [], None)
        atual = janela._modelo_atual()
        janela.painel.documento_corpo().setModified(True)
        ok = janela._gravar(pasta / "rodape.docx", trocar=False)
        janela.painel.sair_da_faixa()
        janela.painel.documento_corpo().setModified(False)
        janela.close()
        return ok, _estilos_dos_campos(atual)

    resultado_certo = campos.PROP_CAMPO_RESULTADO
    try:
        with tempfile.TemporaryDirectory(prefix="editdoc-campo-") as nome:
            pasta = pathlib.Path(nome)
            ok, vindos = salvar_pela_janela(pasta)
            checa_igual(vindos, esperado, "da tela de volta ao modelo, o campo guarda o estilo")
            checa(ok and not falhas, "o arquivo com campo em estilo de caractere grava")
            relido = ler((pasta / "rodape.docx").read_bytes())[0]
            checa_igual(_estilos_dos_campos(relido), esperado, "o estilo do campo esta' no arquivo")

            # Contraprova: com o numero repetido da 0.1.10, o resultado "1" vira o estilo "1"
            # e a conferencia recusa o arquivo (o defeito do documento real).
            campos.PROP_CAMPO_RESULTADO = pq_estilos.PROP_ESTILO_CAR
            (pasta / "rodape.docx").unlink()
            ok, vindos = salvar_pela_janela(pasta)
            checa(("PAGE", "1") in vindos, "contraprova: o resultado vira o estilo de id \"1\"")
            checa(not ok and falhas == ["rodape.docx"] and not (pasta / "rodape.docx").exists(),
                  "contraprova: a conferencia recusa e nada e' gravado")
    finally:
        campos.PROP_CAMPO_RESULTADO = resultado_certo
        dialogos_modais.avisar, dialogos_modais.pedir_falha_na_conferencia = originais

    # Contraprova da conferencia (0.1.11): o campo que perde o estilo, com outro run do mesmo
    # estilo no paragrafo, era aceito; agora e' acusado.
    sem_estilo = ler(dados)[0]
    for r in sem_estilo.paragrafos[1].runs:
        if r.campo:
            r.estilo = None
    sem_estilo.paragrafos[1].runs.append(Run(" fim", estilo="Nmerodepgina"))
    com_estilo = ler(dados)[0]
    com_estilo.paragrafos[1].runs.append(Run(" fim", estilo="Nmerodepgina"))
    try:
        conferencia.conferir(escrever(sem_estilo), com_estilo)
        checa(False, "contraprova: o campo sem estilo passaria pela conferencia")
    except conferencia.FalhaDeConferencia as erro:
        checa("estilo" in str(erro), f"contraprova: o campo que perdeu o estilo e' acusado ({erro})")


def testar_propriedades_unicas() -> None:
    secao("propriedades de usuario do Qt: um numero para cada uma (0.1.11)")
    import importlib
    import pkgutil

    import editdoc.ponte_qt as pacote_ponte

    def repetidas(valores: dict[str, set[int]]) -> list:
        # PROP_ESTILO (bloco) e PROP_CAMPO (caractere) dividem o + 4 sem colidir: formatos
        # diferentes. Qualquer outra repeticao entre nomes diferentes e' defeito.
        por_numero: dict[int, set[str]] = {}
        for nome, numeros in valores.items():
            for n in numeros:
                por_numero.setdefault(n, set()).add(nome)
        return sorted((n, sorted(nomes)) for n, nomes in por_numero.items()
                      if len(nomes) > 1 and nomes != {"PROP_ESTILO", "PROP_CAMPO"})

    valores: dict[str, set[int]] = {}
    for info in pkgutil.iter_modules(pacote_ponte.__path__):
        modulo = importlib.import_module(f"editdoc.ponte_qt.{info.name}")
        for nome, valor in vars(modulo).items():
            if nome.lstrip("_").startswith("PROP_") and isinstance(valor, int):
                valores.setdefault(nome.lstrip("_"), set()).add(valor)
    checa(len(valores) >= 20, f"as propriedades foram achadas ({len(valores)})")
    checa_igual(repetidas(valores), [], "nenhum numero de propriedade repetido")
    sabotado = dict(valores)
    sabotado["PROP_CAMPO_RESULTADO"] = valores["PROP_ESTILO_CAR"]
    checa(repetidas(sabotado) != [], "contraprova: o numero repetido da 0.1.10 e' acusado")


def testar_valor_na_tela() -> None:
    secao("tela: o campo do corpo mostra o numero da sua pagina, e o total acompanha")
    from editdoc.interface.painel_documento import PainelDocumento
    from editdoc.ponte_qt import campos, para_qt

    vistos: list[tuple[str, str]] = []
    original = campos.ManipuladorDeCampo.drawObject

    def espiao(self, painter, rect, doc, pos, fmt):
        instrucao = campos.instrucao_de(fmt)
        vistos.append((instrucao, self.contexto.valor(instrucao, self._indice(rect.center().y()))))
        return original(self, painter, rect, doc, pos, fmt)

    campos.ManipuladorDeCampo.drawObject = espiao
    try:
        blocos = []
        for pagina in range(4):
            blocos.append(Paragrafo(runs=[Run("Esta e' a pagina "), Run("1", campo="PAGE"),
                                          Run(" de "), Run("1", campo="NUMPAGES")],
                                    quebra_antes=pagina > 0))
        painel = PainelDocumento("Calibri", 11)
        painel.resize(900, 1200)
        painel.definir_zoom(0.25)
        painel.show()
        para_qt.preencher(painel.document(), Documento(paragrafos=blocos))
        painel._limpar_cache()
        drenar_eventos()
        painel.viewport().grab()
        paginas = [v for i, v in vistos if i == "PAGE"]
        totais = {v for i, v in vistos if i == "NUMPAGES"}
        checa_igual(sorted(set(paginas)), ["1", "2", "3", "4"], "PAGE mostra 1, 2, 3 e 4, um por pagina")
        checa_igual(totais, {"4"}, "NUMPAGES mostra 4 em todas")
        # Mais uma pagina: o total muda sem editar o campo.
        cursor = painel.textCursor()
        cursor.movePosition(cursor.MoveOperation.End)
        painel.setTextCursor(cursor)
        painel.inserir_quebra_de_pagina()
        vistos.clear()
        drenar_eventos()
        painel.viewport().grab()
        checa_igual({v for i, v in vistos if i == "NUMPAGES"}, {"5"},
                    "com uma pagina a mais, NUMPAGES passa a 5")
        painel.definir_configuracao_pagina(Pagina(formato_numero="upperRoman"))
        vistos.clear()
        drenar_eventos()
        painel.viewport().grab()
        checa("IV" in {v for i, v in vistos if i == "PAGE"}, "formato romano: a 4a pagina mostra IV")
        painel.document().setModified(False)
        painel.close()
    finally:
        campos.ManipuladorDeCampo.drawObject = original


def main() -> int:
    testar_numeros()
    testar_corpo_empurrado()
    testar_leitura()
    testar_escrita()
    testar_oraculo()
    if not QT:
        return pular("PySide6 nao esta instalado neste venv")
    if not usar_fonte_windows(9):
        return pular("Segoe UI nao esta instalada")
    testar_ponte_qt()
    testar_estilo_do_campo()
    testar_propriedades_unicas()
    testar_valor_na_tela()
    return resumir()


if __name__ == "__main__":
    raise SystemExit(main())
