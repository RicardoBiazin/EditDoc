"""python-docx como oraculo independente (so' leitura): texto, formato e paragrafo.

O escritor e o leitor sao nossos; o python-docx foi escrito por outra pessoa e le
o que o Word leria. Divergencia aqui e' um defeito do nosso XML, nao do teste.
A contraprova troca um negrito no XML e a divergencia tem de aparecer.
"""

from __future__ import annotations

import io
import re

from ajudantes import checa, checa_igual, resumir, secao
import fixtures_docx as fx

import docx  # python-docx 1.2.0: so' no ambiente de desenvolvimento
from docx.enum.text import WD_ALIGN_PARAGRAPH

from editdoc.modelo.documento import Documento

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_ALINH = {"esquerda": WD_ALIGN_PARAGRAPH.LEFT, "centro": WD_ALIGN_PARAGRAPH.CENTER,
          "direita": WD_ALIGN_PARAGRAPH.RIGHT, "justificado": WD_ALIGN_PARAGRAPH.JUSTIFY}


def _numfmt(documento, num_id: str) -> str | None:
    """Formato do nivel 0 da lista, lido pelo python-docx a partir do numId do paragrafo."""
    numeracao = documento.part.numbering_part.element
    for n in numeracao.findall(_W + "num"):
        if n.get(_W + "numId") == num_id:
            aid = n.find(_W + "abstractNumId").get(_W + "val")
            for a in numeracao.findall(_W + "abstractNum"):
                if a.get(_W + "abstractNumId") == aid:
                    return a.find(".//" + _W + "numFmt").get(_W + "val")
    return None


def divergencias(modelo: Documento, dados: bytes) -> list[str]:
    """Tudo o que o python-docx le de forma diferente do modelo. Lista vazia = concorda."""
    d = docx.Document(io.BytesIO(dados))
    saida: list[str] = []
    paras = d.paragraphs
    if len(paras) != len(modelo.paragrafos):
        return [f"paragrafos: {len(paras)} contra {len(modelo.paragrafos)}"]
    for i, (mp, dp) in enumerate(zip(modelo.paragrafos, paras)):
        runs_m = [r for r in mp.runs if r.texto]
        texto_m = "".join(r.texto for r in runs_m)
        if dp.text != texto_m:
            saida.append(f"p{i}: texto {dp.text!r} contra {texto_m!r}")
        if (dp.alignment is None) != (mp.alinhamento is None) or (
                mp.alinhamento and dp.alignment != _ALINH[mp.alinhamento]):
            saida.append(f"p{i}: alinhamento")
        pf = dp.paragraph_format
        for nome, esperado, obtido in (
                ("recuo esquerdo", mp.recuo_esquerdo_tw, pf.left_indent),
                ("recuo da primeira linha", mp.recuo_primeira_linha_tw, pf.first_line_indent),
                ("espaco antes", mp.espaco_antes_tw, pf.space_before),
                ("espaco depois", mp.espaco_depois_tw, pf.space_after)):
            valor = None if obtido is None else obtido.twips
            if valor != esperado:
                saida.append(f"p{i}: {nome} {valor} contra {esperado}")
        if mp.espacamento_linha_pct is not None:
            if abs((pf.line_spacing or 0) - mp.espacamento_linha_pct / 100) > 1e-6:
                saida.append(f"p{i}: espacamento entre linhas")
        if mp.numeracao is not None:
            ids = dp._p.xpath("./w:pPr/w:numPr/w:numId/@w:val")
            fmt = _numfmt(d, ids[0]) if ids else None
            esperado = "bullet" if mp.numeracao == "marcadores" else "decimal"
            if fmt != esperado:
                saida.append(f"p{i}: lista {fmt!r} contra {esperado!r}")
        elif dp._p.xpath("./w:pPr/w:numPr"):
            saida.append(f"p{i}: lista sobrando")
        for j, (rm, rd) in enumerate(zip(runs_m, dp.runs)):
            tag = f"p{i}r{j}"
            for nome, esperado, obtido in (
                    ("negrito", rm.negrito, rd.bold),
                    ("italico", rm.italico, rd.italic),
                    ("sublinhado", rm.sublinhado, rd.underline),
                    ("tachado", rm.tachado, rd.font.strike)):
                if esperado is None:
                    if obtido is not None:
                        saida.append(f"{tag}: {nome} deveria herdar")
                elif bool(obtido) != esperado:
                    saida.append(f"{tag}: {nome} {obtido} contra {esperado}")
            if bool(rd.font.subscript) != bool(rm.subscrito):
                saida.append(f"{tag}: subscrito")
            if bool(rd.font.superscript) != bool(rm.sobrescrito):
                saida.append(f"{tag}: sobrescrito")
            tamanho = None if rd.font.size is None else rd.font.size.pt * 2
            if tamanho != rm.tamanho_meio_pt and not (tamanho is None and rm.tamanho_meio_pt is None):
                saida.append(f"{tag}: tamanho {tamanho} contra {rm.tamanho_meio_pt}")
            if rm.fonte and rd.font.name != rm.fonte:
                saida.append(f"{tag}: fonte {rd.font.name!r}")
            if rm.cor_hex and str(rd.font.color.rgb) != rm.cor_hex:
                saida.append(f"{tag}: cor {rd.font.color.rgb}")
        if len(runs_m) != len(dp.runs):
            saida.append(f"p{i}: quantidade de runs {len(dp.runs)} contra {len(runs_m)}")
    return saida


def testar_oraculo() -> None:
    secao("python-docx le o que escrevemos: texto, formato, paragrafo e listas")
    modelo = fx.documento_completo()
    dados = fx.bytes_de(modelo)
    d = docx.Document(io.BytesIO(dados))
    checa_igual(d.core_properties.title, "Amostra do EditDoc", "titulo do documento")
    checa_igual(d.core_properties.author, "Teste", "autor do documento")
    achados = divergencias(modelo, dados)
    checa_igual(achados, [], "nenhuma divergencia entre o modelo e o python-docx")
    misturado = fx.documento_com_formatos_mistos()
    checa_igual(divergencias(misturado, fx.bytes_de(misturado)), [],
                "runs com None, False e True: python-docx concorda")


def testar_contraprova() -> None:
    secao("contraprova: um negrito trocado no XML e' acusado pelo oraculo")
    modelo = fx.documento_completo()
    dados = fx.bytes_de(modelo)
    xml = fx.partes_de(dados)["word/document.xml"]
    trocado = re.sub(rb"<w:b/>", b'<w:b w:val="0"/>', xml, count=1)
    checa(trocado != xml, "a sabotagem mudou o XML")
    adulterado = fx.com_parte(dados, "word/document.xml", trocado)
    achados = divergencias(modelo, adulterado)
    checa(len(achados) >= 1, f"o negrito desligado e' acusado ({len(achados)} divergencia(s))")


def main() -> int:
    testar_oraculo()
    testar_contraprova()
    return resumir()


if __name__ == "__main__":
    raise SystemExit(main())
