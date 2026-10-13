"""Estilos (M6): catalogo, heranca, styles.xml do Word ida e volta, ponte Qt e comandos.

Cada checagem que pode passar por acaso tem contraprova: o estilo que o editor nao usa
volta byte a byte (e o alterado nao); a referencia quebrada e' acusada pelo validador;
o desfazer devolve o catalogo E os paragrafos.
"""

from __future__ import annotations

import copy
import io
import pathlib
import zipfile

from ajudantes import (RAIZ, checa, checa_igual, drenar_eventos, pasta_temporaria, preparar_qt,
                       pular, resumir, secao, usar_fonte_windows)

if not preparar_qt():
    raise SystemExit(pular("PySide6 nao esta instalado neste venv"))

from lxml import etree  # noqa: E402
from PySide6.QtCore import Qt  # noqa: E402
from PySide6.QtGui import QColor, QFont, QTextCharFormat, QTextCursor, QTextDocument  # noqa: E402
from PySide6.QtTest import QTest  # noqa: E402

import fixtures_docx as fx  # noqa: E402
from editdoc import configuracao  # noqa: E402
from editdoc.edicao import estilos as ed  # noqa: E402
from editdoc.edicao import formatacao  # noqa: E402
from editdoc.edicao.estado import ler  # noqa: E402
from editdoc.formato_docx import ordem  # noqa: E402
from editdoc.formato_docx.escritor import escrever, escrever_com_avisos  # noqa: E402
from editdoc.formato_docx.leitor import ler as ler_docx  # noqa: E402
from editdoc.formato_docx.nomes import W, q  # noqa: E402
from editdoc.interface import dialogo_estilo  # noqa: E402
from editdoc.interface.dialogo_estilo import DialogoEstilo  # noqa: E402
from editdoc.interface.janela_principal import JanelaPrincipal  # noqa: E402
from editdoc.interface.painel_documento import PainelDocumento  # noqa: E402
from editdoc.modelo.documento import Documento, Paragrafo, Run, Tabela, Linha, Celula, normalizar  # noqa: E402
from editdoc.modelo.estilos import Estilo, Estilos, catalogo_padrao  # noqa: E402
from editdoc.ponte_qt import de_qt, para_qt  # noqa: E402
from editdoc.ponte_qt import estilos as pq  # noqa: E402

CTRL = Qt.KeyboardModifier.ControlModifier
ALT = Qt.KeyboardModifier.AltModifier
SHIFT = Qt.KeyboardModifier.ShiftModifier


def _partes(dados: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(dados)) as z:
        return {n: z.read(n) for n in z.namelist()}


def _estilo_xml(raiz, sid):
    return next((s for s in raiz.findall(q("style")) if s.get(q("styleId")) == sid), None)


def _c14n(el) -> bytes:
    return etree.tostring(el, method="c14n")


# -- modelo -------------------------------------------------------------------------

def testar_catalogo_padrao() -> None:
    secao("conjunto padrao do documento novo: ids do Word, nomes em portugues")
    cat = catalogo_padrao()
    ids = set(cat.itens)
    esperados = {"Normal", "Title", "Subtitle", "Heading1", "Heading2", "Heading3", "Quote",
                 "Emphasis", "Strong", "ListParagraph", "NoSpacing"}
    checa(esperados <= ids, f"o documento novo tem os estilos do Word ({sorted(esperados - ids)} faltando)")
    nomes = {e.id: e.nome_exibido() for e in cat.itens.values()}
    checa_igual([nomes[i] for i in ("Title", "Subtitle", "Heading1", "Heading3", "Quote", "Emphasis",
                                     "Strong", "ListParagraph", "NoSpacing")],
                ["Título", "Subtítulo", "Título 1", "Título 3", "Citação", "Ênfase", "Forte",
                 "Parágrafo da Lista", "Sem Espaçamento"], "nomes na tela em portugues, com acento")
    checa_igual((cat.obter("Emphasis").tipo, cat.obter("Strong").tipo), ("caractere", "caractere"),
                "Ênfase e Forte sao estilos de caractere")
    checa_igual(cat.obter("Heading1").proximo, "Normal", "depois do Titulo 1 vem o Normal")
    checa_igual(cat.resolver_paragrafo("Heading2")[0].get("nivel"), 1, "Titulo 2 tem nivel de estrutura 1")
    checa_igual(cat.id_padrao(), "Normal", "Normal e' o estilo padrao")


def testar_heranca() -> None:
    secao("heranca basedOn: resolucao, ciclo e estilo interno criado ao usar")
    cat = Estilos(itens={
        "A": Estilo("A", "A", caractere={"tamanho_meio_pt": 30, "negrito": True},
                    paragrafo={"alinhamento": "centro"}),
        "B": Estilo("B", "B", baseado_em="A", caractere={"cor_hex": "FF0000"}),
        "C": Estilo("C", "C", baseado_em="B", caractere={"tamanho_meio_pt": 20},
                    paragrafo={"espaco_antes_tw": 120}),
        "X": Estilo("X", "X", baseado_em="Y", caractere={"italico": True}),
        "Y": Estilo("Y", "Y", baseado_em="X", caractere={"negrito": True}),
    })
    par, car = cat.resolver_paragrafo("C")
    checa_igual(car, {"tamanho_meio_pt": 20, "negrito": True, "cor_hex": "FF0000"},
                "C herda de B e de A; o tamanho proprio vence")
    checa_igual(par, {"alinhamento": "centro", "espaco_antes_tw": 120}, "paragrafo tambem herda")
    checa_igual(cat.resolver_paragrafo("B")[1].get("tamanho_meio_pt"), 30,
                "contraprova: B nao ve o tamanho de C, que herda dele")
    checa_igual(cat.resolver_paragrafo("X")[1], {"italico": True, "negrito": True},
                "ciclo X <-> Y para sem travar")
    novo = cat.copia()
    criado = novo.garantir("heading 2")
    checa(criado is not None and criado.id == "Heading2" and "Heading2" not in cat.itens,
          "garantir cria o Titulo 2 que faltava (sem mexer no catalogo original)")
    volta = Estilos.de_json(novo.para_json(), novo)
    checa_igual(volta.itens, novo.itens, "o catalogo vai e volta pelo JSON (desfazer da tela)")


# -- styles.xml do Word ----------------------------------------------------------------

def testar_word_ida_e_volta() -> None:
    secao("styles.xml do Word: leitura, ida e volta sem perder estilo, tema e latentStyles")
    dados = fx.docx_com_estilos_do_word()
    doc, avisos = ler_docx(dados)
    cat = doc.estilos
    ids_orig = {s.get(q("styleId")) for s in etree.fromstring(fx.ESTILOS_DO_WORD.encode()).findall(q("style"))}
    checa_igual(set(cat.itens), ids_orig, "todos os estilos do arquivo entram no catalogo")
    estilos_p = [p.estilo for p in doc.paragrafos]
    checa_igual(estilos_p, ["Ttulo1", None, "Ttulo2", "Commarcadores", None, None],
                "pStyle de cada paragrafo (Normal explicito vira o padrao)")
    checa_igual(cat.obter("Ttulo1").nome_exibido(), "Título 1",
                "o id em portugues sem acento mostra o nome pelo nome canonico")
    _par, car = cat.resolver_paragrafo("Ttulo1")
    checa_igual(car.get("fonte"), "Cambria", "fonte de tema (majorHAnsi) resolvida pelo theme1.xml")
    checa_igual(cat.padrao_caractere, {"fonte": "Garamond", "tamanho_meio_pt": 24},
                "docDefaults: fonte de tema minorHAnsi e 12 pt")
    run_destaque = doc.paragrafos[1].runs[1]
    checa_igual((run_destaque.estilo, run_destaque.negrito), ("Destaque", None),
                "rStyle vira Run.estilo; o negrito do estilo nao vira formatacao direta")
    checa_igual(doc.paragrafos[2].runs[1].cor_hex, "FF0000", "a cor direta dentro do Titulo 2 fica direta")
    checa_igual(doc.paragrafos[3].numeracao, "marcadores", "a lista do estilo (numPr) passa ao paragrafo")
    checa_igual(doc.paragrafos[4].nivel, 0, "outlineLvl direto no paragrafo e' lido")
    codigos = {a.codigo: a.contagem for a in avisos}
    checa_igual(codigos, {"FORMATACAO": 1}, "so' o rPrChange do estilo gera aviso (nada de ESTILO)")

    gravado = escrever(doc)
    partes = _partes(gravado)
    raiz = etree.fromstring(partes["word/styles.xml"])
    original = etree.fromstring(fx.ESTILOS_DO_WORD.encode())
    ids_gravados = {s.get(q("styleId")) for s in raiz.findall(q("style"))}
    checa(ids_orig <= ids_gravados, f"nenhum estilo do arquivo se perde ({sorted(ids_orig - ids_gravados)})")
    for sid in ("MeuEstiloRaro", "Ttulo1", "Destaque", "TabeladeGrade1Clara", "Ttulo1Char"):
        checa_igual(_c14n(_estilo_xml(raiz, sid)), _c14n(_estilo_xml(original, sid)),
                    f"estilo {sid} volta igual, byte a byte (c14n)")
    checa_igual(len(raiz.find(q("latentStyles"))), 4, "latentStyles volta inteiro")
    checa(raiz.get("{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable") == "w14 w15",
          "mc:Ignorable e as extensoes w14/w15 continuam declarados")
    checa_igual(_c14n(raiz.find(q("docDefaults"))), _c14n(original.find(q("docDefaults"))),
                "docDefaults volta igual (com w14:ligatures)")
    checa(partes.get("word/theme/theme1.xml") == fx.TEMA_DO_WORD.encode(), "o tema volta com os mesmos bytes")
    # Reescritas na 0.1.4 (decisao do M6 revertida): o numbering.xml de origem volta
    # inteiro, entao o numPr fica no estilo e o paragrafo herda, sem numPr direto.
    checa(_estilo_xml(raiz, "Commarcadores").find(q("pPr") + "/" + q("numPr")) is not None,
          "o numPr fica no estilo (o numbering de origem volta)")
    checa(b"<w:numPr>" not in partes["word/document.xml"].split(b"item da lista")[0].rsplit(b"<w:p>", 1)[1],
          "...e o paragrafo da lista herda do estilo, sem numPr direto")
    checa(_estilo_xml(raiz, "ComRevisao").find(".//" + q("rPrChange")) is None,
          "o rPrChange do estilo sai (o aviso FORMATACAO disse)")
    checa_igual(ordem.validar(partes), [], "o pacote gravado passa no validador")

    relido, avisos2 = ler_docx(gravado)
    checa(relido == doc, "segunda leitura: mesmo modelo (texto, formatacao direta e estilos)")
    checa_igual([a.codigo for a in avisos2], [], "segunda leitura sem avisos")
    checa_igual(relido.estilos.resolver_paragrafo("Ttulo1"), cat.resolver_paragrafo("Ttulo1"),
                "Titulo 1 resolve igual depois da ida e volta")
    checa_igual(escrever(relido), escrever(ler_docx(escrever(relido))[0]),
                "idempotente: gravar o relido da' os mesmos bytes")

    from docx import Document as DocxOraculo
    oraculo = DocxOraculo(io.BytesIO(gravado))
    checa_igual([p.style.name for p in oraculo.paragraphs][:3], ["Heading 1", "Normal", "Heading 2"],
                "python-docx ve os estilos de paragrafo")


def testar_estilo_alterado_no_xml() -> None:
    secao("estilo alterado: so' o que mudou e' reescrito")
    doc, _ = ler_docx(fx.docx_com_estilos_do_word())
    e = doc.estilos.obter("Ttulo1")
    e.caractere["cor_hex"] = "00B050"
    e.paragrafo["alinhamento"] = "centro"
    e.alterado = True
    raiz = etree.fromstring(_partes(escrever(doc))["word/styles.xml"])
    el = _estilo_xml(raiz, "Ttulo1")
    cor = el.find(q("rPr") + "/" + q("color"))
    checa(cor is not None and cor.get(q("val")) == "00B050" and cor.get(q("themeColor")) is None,
          "cor nova gravada, sem o themeColor que venceria o valor")
    checa(el.find(q("pPr") + "/" + q("jc")).get(q("val")) == "center", "alinhamento novo gravado")
    checa(el.find(q("pPr") + "/" + q("keepNext")) is not None, "keepNext, que o modelo nao conhece, fica")
    checa(el.find(q("rPr") + "/" + q("rFonts")).get(q("asciiTheme")) == "majorHAnsi",
          "a fonte de tema, que nao mudou, fica de tema")
    checa_igual(el.find(q("rPr") + "/" + q("sz")).get(q("val")), "40", "o tamanho que nao mudou fica")
    checa(el.find(q("rsid")) is not None, "o rsid do Word fica")
    relido, _ = ler_docx(escrever(doc))
    checa_igual(relido.estilos.resolver_paragrafo("Ttulo1")[1].get("cor_hex"), "00B050",
                "relido: a cor nova vale")


def testar_validador_de_referencias() -> None:
    secao("validador: pStyle e rStyle apontam para estilo que existe, do tipo certo")
    doc = Documento(paragrafos=[Paragrafo(runs=[Run("x", estilo="Strong")], estilo="Heading1")])
    partes = _partes(escrever(doc))
    checa_igual(ordem.validar_estilos(partes), [], "referencias validas passam")
    sabotado = dict(partes)
    sabotado["word/document.xml"] = partes["word/document.xml"].replace(b'w:val="Heading1"', b'w:val="NaoExiste"')
    checa(ordem.validar_estilos(sabotado), "contraprova: pStyle para estilo inexistente e' acusado")
    sabotado["word/document.xml"] = partes["word/document.xml"].replace(b'w:val="Strong"', b'w:val="Heading1"')
    checa(ordem.validar_estilos(sabotado), "contraprova: rStyle para estilo de paragrafo e' acusado")
    fantasma = Documento(paragrafos=[Paragrafo(runs=[Run("x")], estilo="NaoTem")])
    xml = _partes(escrever(fantasma))["word/document.xml"]
    checa(b"NaoTem" not in xml, "estilo fora do catalogo nao e' gravado (o paragrafo fica Normal)")


# -- ponte Qt ---------------------------------------------------------------------------

def testar_ponte_qt() -> None:
    secao("ponte Qt: o Qt recebe o efetivo; a volta e' so' a formatacao direta")
    doc, _ = ler_docx(fx.docx_com_estilos_do_word())
    qdoc = QTextDocument()
    para_qt.preencher(qdoc, doc)
    titulo = qdoc.firstBlock()
    cf = QTextCursor(titulo)
    cf.movePosition(QTextCursor.MoveOperation.Right)
    checa_igual(cf.charFormat().fontPointSize(), 20.0, "o Titulo 1 aparece com 20 pt (o do estilo)")
    checa_igual(pq.nivel_do_bloco(titulo), 0, "o bloco do titulo tem nivel 0")
    checa_igual(pq.estilo_do_bloco(titulo.blockFormat()), "Ttulo1", "o id do estilo fica no bloco")
    checa_igual(pq.nivel_do_bloco(titulo.next()), None, "contraprova: paragrafo Normal sem nivel")
    volta = de_qt.extrair(qdoc)
    # 0.1.4: o documento de comparacao leva os catalogos (estilos e numeracao): sem eles o tipo
    # da lista herdada do estilo pareceria uma lista a criar.
    checa(normalizar(volta) == normalizar(Documento(paragrafos=doc.paragrafos, estilos=doc.estilos,
                                                    numeracao=doc.numeracao)),
          "Qt -> modelo devolve o mesmo modelo (estilos e formatacao direta)")
    checa_igual(volta.paragrafos[0].runs[0].tamanho_meio_pt, None,
                "o tamanho do estilo nao vira formatacao direta na volta")


# -- comandos -------------------------------------------------------------------------------

def _painel(texto: str) -> PainelDocumento:
    p = PainelDocumento("Calibri", 11)
    p.setPlainText(texto)
    p.document().clearUndoRedoStacks()
    return p


def _por(p, inicio, fim=None) -> None:
    c = p.textCursor()
    c.setPosition(inicio)
    if fim is not None:
        c.setPosition(fim, QTextCursor.MoveMode.KeepAnchor)
    p.setTextCursor(c)


def _um_passo(p, comando, mensagem) -> None:
    antes = p.document().toHtml()
    cat_antes = pq.catalogo(p.document()).para_json()
    c = p.textCursor()
    selecao = (c.anchor(), c.position()) if c.hasSelection() else (c.position(),)
    comando()
    mudou = p.document().toHtml() != antes
    p.undo()
    voltou = p.document().toHtml() == antes and pq.catalogo(p.document()).para_json() == cat_antes
    # O desfazer move o cursor (como no QTextEdit): a selecao do teste volta para o proximo passo.
    _por(p, *selecao)
    checa(mudou and voltou, f"{mensagem} (mudou: {mudou}, um desfazer voltou tudo: {voltou})")


def _modelo(p):
    return [(x.estilo, [(r.texto, r.estilo, r.tamanho_meio_pt, r.cor_hex, r.negrito, r.italico)
                        for r in x.runs]) for x in de_qt.extrair(p.document()).paragrafos]


def testar_aplicar_e_limpar() -> None:
    secao("aplicar estilo de paragrafo e de caractere; limpar")
    p = _painel("Um titulo\nTexto comum aqui")
    _por(p, 2)
    _um_passo(p, lambda: ed.aplicar_estilo(p, "heading 1"), "aplicar Titulo 1 e' um passo de desfazer")
    ed.aplicar_estilo(p, "heading 1")
    checa_igual((ed.estilo_no_cursor(p)[0], ler(p).tamanho), ("Heading1", 16.0),
                "Titulo 1 aplicado: 16 pt no cursor")
    _por(p, 10, 15)
    _um_passo(p, lambda: ed.aplicar_estilo(p, "Strong"), "aplicar Forte e' um passo de desfazer")
    ed.aplicar_estilo(p, "Strong")
    m = _modelo(p)
    checa_igual(m[1][1][0][:2], ("Texto", "Strong"), "Forte vira estilo de caractere do trecho")
    checa(ler(p).negrito and m[1][1][0][4] is None, "o trecho fica em negrito, sem negrito direto")
    # Limpar formatacao de caractere (Ctrl+Espaco): tira o estilo de caractere e a direta.
    _por(p, 10, 15)
    formatacao.alternar_italico(p)
    formatacao.limpar_caractere(p)
    m = _modelo(p)
    checa(all(r[1] is None and r[5] is None for r in m[1][1]), "limpar caractere tira Forte e o italico")
    _por(p, 0, 15)
    _um_passo(p, lambda: formatacao.limpar_formatacao(p), "limpar toda a formatacao e' um passo")
    formatacao.limpar_formatacao(p)
    checa_igual([x[0] for x in _modelo(p)], [None, None], "limpar toda a formatacao volta ao Normal")
    checa_igual(ler(p).tamanho, 11.0, "e o tamanho volta ao padrao")


def testar_lista_com_estilo() -> None:
    secao("lista convive com estilo: aplicar Titulo numa lista mantem a lista")
    from PySide6.QtGui import QTextListFormat
    p = _painel("item um\nitem dois")
    _por(p, 0, 12)
    formatacao.alternar_lista(p, QTextListFormat.Style.ListDisc)
    _por(p, 0)
    ed.aplicar_estilo(p, "heading 2")
    bloco = p.document().firstBlock()
    from editdoc.edicao import listas as ed_listas
    checa(ed_listas.eh_item(p, bloco), "o paragrafo continua na lista")  # 0.1.4: sem QTextList
    m = de_qt.extrair(p.document()).paragrafos[0]
    checa_igual((m.estilo, m.numeracao), ("Heading2", "marcadores"), "modelo: estilo e lista juntos")


def testar_enter_proximo() -> None:
    secao("Enter no fim do titulo: o paragrafo seguinte usa o estilo next")
    usar_fonte_windows(9)
    p = _painel("Capitulo\ncorpo")
    _por(p, 0)
    ed.aplicar_estilo(p, "heading 1")
    _por(p, len("Capitulo"))
    QTest.keyClick(p, Qt.Key.Key_Return)
    QTest.keyClicks(p, "novo")
    m = _modelo(p)
    checa_igual([x[0] for x in m], ["Heading1", None, None], "o paragrafo novo e' Normal")
    checa_igual(m[1][1][0][2], None, "e o texto digitado nele nao herda os 16 pt")
    checa(ler(p).tamanho == 11.0, "a faixa mostra 11 pt no paragrafo novo")
    p.undo()
    p.undo()
    _por(p, 4)
    QTest.keyClick(p, Qt.Key.Key_Return)
    checa_igual([x[0] for x in _modelo(p)][:2], ["Heading1", "Heading1"],
                "contraprova: Enter no meio do titulo deixa as duas metades no Titulo 1")


def testar_modificar_estilo() -> None:
    secao("modificar estilo: todos os paragrafos mudam, a formatacao direta fica, um desfazer")
    p = _painel("Titulo A\nmeio\nTitulo B")
    for pos in (0, 15):
        _por(p, pos)
        ed.aplicar_estilo(p, "heading 1")
    _por(p, 0, 6)
    fmt = QTextCharFormat()
    fmt.setForeground(QColor("#FF0000"))
    p.textCursor().mergeCharFormat(fmt)
    p.document().clearUndoRedoStacks()
    antes = p.document().toHtml()
    novo = copy.deepcopy(ed.catalogo_de(p).obter("Heading1"))
    novo.caractere.update({"tamanho_meio_pt": 40, "cor_hex": "00B050"})
    ed.redefinir_estilo(p, novo)
    m = _modelo(p)
    checa_igual(m[0][1][0][3], "FF0000", "o trecho com cor direta continua vermelho")
    checa(all(r[2] is None for r in m[0][1]) and m[2][1][0][3] is None,
          "o resto dos dois titulos segue o estilo (sem direta)")
    c = QTextCursor(p.document().lastBlock())
    c.movePosition(QTextCursor.MoveOperation.Right)
    checa_igual((c.charFormat().fontPointSize(), c.charFormat().foreground().color().name()),
                (20.0, "#00b050"), "o Titulo B na tela: 20 pt e verde")
    meio = QTextCursor(p.document().findBlockByNumber(1))
    meio.movePosition(QTextCursor.MoveOperation.Right)
    checa(not meio.charFormat().hasProperty(QTextCharFormat.Property.FontPointSize),
          "contraprova: o paragrafo Normal nao muda")
    p.undo()
    checa(p.document().toHtml() == antes, "um desfazer devolve os paragrafos")
    checa_igual(ed.catalogo_de(p).obter("Heading1").caractere.get("tamanho_meio_pt"), 32,
                "...e o catalogo (o estilo volta a 16 pt)")
    p.redo()
    checa_igual(ed.catalogo_de(p).obter("Heading1").caractere.get("cor_hex"), "00B050",
                "refazer volta a definicao nova")


def testar_atualizar_e_criar() -> None:
    secao("atualizar estilo pela selecao e criar estilo novo")
    p = _painel("Secao um\nSecao dois\ntexto")
    for pos in (0, 9):
        _por(p, pos)
        ed.aplicar_estilo(p, "heading 2")
    _por(p, 0, 8)
    formatacao.alternar_italico(p)
    formatacao.alinhar(p, Qt.AlignmentFlag.AlignHCenter)
    _por(p, 2)
    _um_passo(p, lambda: ed.atualizar_pela_selecao(p), "atualizar o estilo e' um passo de desfazer")
    ed.atualizar_pela_selecao(p)
    est = ed.catalogo_de(p).obter("Heading2")
    checa_igual((est.caractere.get("italico"), est.paragrafo.get("alinhamento")), (True, "centro"),
                "Titulo 2 passa a ser italico e centralizado")
    m = de_qt.extrair(p.document()).paragrafos
    checa(m[1].alinhamento is None and m[1].runs[0].italico is None and
          m[1].runs[0].texto == "Secao dois" and
          p.document().findBlockByNumber(1).blockFormat().alignment() & Qt.AlignmentFlag.AlignHCenter,
          "o outro Titulo 2 ficou centralizado, sem formatacao direta")
    checa(m[0].runs[0].italico is None, "no paragrafo de origem a formatacao virou estilo")

    _por(p, 20, 25)
    formatacao.definir_tamanho(p, 18)
    formatacao.alternar_negrito(p)
    _por(p, 22)
    sid = ed.criar_da_selecao(p, "Meu Destaque")
    checa_igual(sid, "MeuDestaque", "id sem espaco nem acento, a partir do nome")
    novo = ed.catalogo_de(p).obter(sid)
    checa_igual((novo.baseado_em, novo.caractere.get("tamanho_meio_pt"), novo.caractere.get("negrito")),
                ("Normal", 36, True), "o estilo novo tem a formatacao da selecao")
    modelo = de_qt.extrair(p.document())
    checa_igual(modelo.paragrafos[2].estilo, sid, "e o paragrafo passa a usa-lo")
    relido, _ = ler_docx(escrever(modelo))
    checa(relido.paragrafos[2].estilo == sid and relido.estilos.obter(sid).nome == "Meu Destaque",
          "o estilo novo vai ao arquivo e volta")
    checa(ed.criar_da_selecao(p, "Meu Destaque") is None, "contraprova: nome repetido nao cria outro")
    p.undo()
    checa(ed.catalogo_de(p).obter(sid) is None, "um desfazer tira o estilo novo do catalogo")


def _janela() -> JanelaPrincipal:
    usar_fonte_windows(9)
    j = JanelaPrincipal(configuracao.padrao())
    j.show()
    drenar_eventos()
    j.painel.setFocus()
    return j


def _fechar(j) -> None:
    j.painel.documento_corpo().setModified(False)
    j.close()


def testar_atalhos() -> None:
    secao("atalhos: Ctrl+Alt+1/2/3 e Ctrl+Shift+N, com teclas reais")
    j = _janela()
    p = j.painel
    QTest.keyClicks(p, "linha")
    for tecla, esperado in ((Qt.Key.Key_1, "Heading1"), (Qt.Key.Key_2, "Heading2"),
                            (Qt.Key.Key_3, "Heading3")):
        QTest.keyClick(p, tecla, CTRL | ALT)
        drenar_eventos()
        checa_igual(ed.estilo_no_cursor(p)[0], esperado, f"Ctrl+Alt+{esperado[-1]} aplica {esperado}")
    QTest.keyClick(p, Qt.Key.Key_N, CTRL | SHIFT)
    checa_igual(ed.estilo_no_cursor(p)[0], "Normal", "Ctrl+Shift+N volta ao Normal")
    QTest.keyClick(p, Qt.Key.Key_1, CTRL)
    checa(ed.estilo_no_cursor(p)[0] == "Normal" and ler(p).espacamento == 100,
          "contraprova: Ctrl+1 continua sendo espacamento simples, nao Titulo 1")
    checa_igual(p.toPlainText(), "linha", "nenhum atalho digitou texto")
    _fechar(j)


def testar_dialogo() -> None:
    secao("dialogo Modificar estilo: so' o que difere da base vira propriedade")
    cat = catalogo_padrao()
    dlg = DialogoEstilo(cat.obter("Heading1"), cat, QFont("Calibri", 11))
    checa_igual(dlg.tamanho.value(), 16.0, "mostra o tamanho do estilo")
    dlg.tamanho.setValue(20)
    dlg.definir_cor("C00000")
    dlg.italico.setChecked(True)
    r = dlg.resultado()
    checa_igual((r.caractere.get("tamanho_meio_pt"), r.caractere.get("cor_hex"), r.caractere.get("italico")),
                (40, "C00000", True), "tamanho, cor e italico novos")
    checa_igual(r.caractere.get("fonte"), "Calibri Light",
                "a fonte que o usuario nao mexeu fica (mesmo sem a fonte instalada)")
    checa_igual(r.paragrafo.get("espaco_antes_tw"), 240, "o espacamento proprio que nao mudou fica")
    checa(dlg.nome.isReadOnly(), "nome de estilo interno nao se edita")
    dlg.deleteLater()
    j = _janela()
    p = j.painel
    QTest.keyClicks(p, "abc")
    QTest.keyClick(p, Qt.Key.Key_1, CTRL | ALT)
    original = dialogo_estilo.pedir_estilo
    try:
        dialogo_estilo.pedir_estilo = lambda _pai, est, _cat, _f, novo=False: r
        checa(j.modificar_estilo(), "janela: Modificar estilo aplica o resultado do dialogo")
        checa_igual(ler(p).tamanho, 20.0, "o texto do Titulo 1 fica com 20 pt")
        criado = copy.deepcopy(ed.modelo_da_selecao(p, "Estilo Teste"))
        criado.caractere["sublinhado"] = True
        dialogo_estilo.pedir_estilo = lambda _pai, est, _cat, _f, novo=False: (
            setattr(est, "caractere", {"sublinhado": True}) or est)
        sid = j.criar_estilo()
        checa(sid is not None and ler(p).sublinhado, "janela: Criar um estilo cria e aplica")
        dialogo_estilo.pedir_estilo = lambda *a, **k: None
        checa(not j.modificar_estilo(), "contraprova: cancelar o dialogo nao muda nada")
    finally:
        dialogo_estilo.pedir_estilo = original
    _fechar(j)


def testar_salvar_e_reabrir() -> None:
    secao("janela: abrir o .docx do Word, mudar estilo, salvar e reabrir")
    j = _janela()
    dados = fx.docx_com_estilos_do_word(estilos_xml=fx.ESTILOS_DO_WORD.replace(
        '<w:rPrChange w:id="1" w:author="Fulano" w:date="2026-01-01T00:00:00Z"><w:rPr/></w:rPrChange>', ""))
    with pasta_temporaria() as pasta:
        alvo = pasta / "estilos.docx"
        alvo.write_bytes(dados)
        modelo, avisos = ler_docx(dados)
        j.carregar(modelo, avisos, alvo)
        fonte = j.painel.documento_corpo().defaultFont()
        checa_igual((fonte.family(), fonte.pointSizeF()), ("Garamond", 12.0),
                    "o docDefaults (tema minor + 12 pt) vira a fonte padrao da tela")
        checa(not j.sujo(), "abrir nao deixa o documento sujo")
        _por(j.painel, 0)
        novo = copy.deepcopy(ed.catalogo_de(j.painel).obter("Ttulo1"))
        novo.caractere["cor_hex"] = "7030A0"
        ed.redefinir_estilo(j.painel, novo)
        checa(j.sujo(), "modificar o estilo deixa o documento sujo")
        checa(j.salvar(), "salvar passa na conferencia")
        relido, _ = ler_docx(alvo.read_bytes())
        checa_igual(relido.estilos.resolver_paragrafo("Ttulo1")[1].get("cor_hex"), "7030A0",
                    "reaberto: o estilo modificado")
        checa("MeuEstiloRaro" in relido.estilos.itens and "TabeladeGrade1Clara" in relido.estilos.itens,
              "reaberto: os estilos que o editor nao usa continuam")
        checa_igual([p.estilo for p in relido.paragrafos], [p.estilo for p in modelo.paragrafos],
                    "reaberto: cada paragrafo com o seu estilo")
    _fechar(j)


def testar_amostra() -> None:
    secao("amostra para abrir no Word: tests/saida/amostra_m6_estilos.docx")
    cat = catalogo_padrao()
    meu = Estilo("Destaque2", "Destaque Azul", baseado_em="Normal", proximo="Normal", rapido=True,
                 personalizado=True, caractere={"cor_hex": "1F4E79", "negrito": True},
                 paragrafo={"recuo_esquerdo_tw": 567, "recuo_direito_tw": 567, "espaco_antes_tw": 120})
    cat.itens[meu.id] = meu
    mod = cat.obter("Heading2")
    mod.caractere["italico"] = True
    p = Paragrafo
    doc = Documento(titulo="Amostra M6", autor="Teste", estilos=cat, paragrafos=[
        p(runs=[Run("Estilos do EditDoc (M6)")], estilo="Title"),
        p(runs=[Run("Subtítulo da amostra")], estilo="Subtitle"),
        p(runs=[Run("1. Título 1")], estilo="Heading1"),
        p(runs=[Run("Parágrafo Normal com "), Run("ênfase", estilo="Emphasis"), Run(" e "),
                Run("forte", estilo="Strong"), Run(" e "), Run("cor direta", cor_hex="C00000"), Run(".")]),
        p(runs=[Run("1.1 Título 2 (modificado: itálico)")], estilo="Heading2"),
        p(runs=[Run("1.1.1 Título 3")], estilo="Heading3"),
        p(runs=[Run("Citação: o estilo Quote, centralizado com recuos.")], estilo="Quote"),
        p(runs=[Run("Estilo criado: Destaque Azul")], estilo="Destaque2"),
        p(runs=[Run("Sem Espaçamento")], estilo="NoSpacing"),
        p(runs=[Run("item com Parágrafo da Lista")], estilo="ListParagraph", numeracao="marcadores"),
        p(runs=[Run("Título 1 com recuo direto")], estilo="Heading1", recuo_esquerdo_tw=720),
        Tabela(colunas_tw=[4513, 4513], linhas=[Linha(celulas=[
            Celula(blocos=[p(runs=[Run("célula com Título 3")], estilo="Heading3")]),
            Celula(blocos=[p(runs=[Run("célula Normal")])])])]),
        p(runs=[Run("fim")]),
    ])
    dados, avisos = escrever_com_avisos(doc)
    saida = pathlib.Path(RAIZ) / "tests" / "saida"
    saida.mkdir(exist_ok=True)
    (saida / "amostra_m6_estilos.docx").write_bytes(dados)
    checa_igual(avisos, [], "a amostra grava sem aviso")
    relido, avisos2 = ler_docx(dados)
    a, b = normalizar(relido).paragrafos, normalizar(doc).paragrafos
    checa_igual([i for i, (x, y) in enumerate(zip(a, b)) if x != y] + [len(a) - len(b)], [0],
                "a amostra relida e' igual ao modelo")
    checa_igual([x.codigo for x in avisos2], [], "e sem avisos na leitura")
    from docx import Document as DocxOraculo
    nomes = [x.style.name for x in DocxOraculo(io.BytesIO(dados)).paragraphs]
    checa("Title" in nomes and "Heading 3" in nomes and "Destaque Azul" in nomes,
          "python-docx abre a amostra e ve os estilos")


def main() -> int:
    testar_catalogo_padrao()
    testar_heranca()
    testar_word_ida_e_volta()
    testar_estilo_alterado_no_xml()
    testar_validador_de_referencias()
    testar_ponte_qt()
    testar_aplicar_e_limpar()
    testar_lista_com_estilo()
    testar_enter_proximo()
    testar_modificar_estilo()
    testar_atualizar_e_criar()
    testar_atalhos()
    testar_dialogo()
    testar_salvar_e_reabrir()
    testar_amostra()
    return resumir()


if __name__ == "__main__":
    raise SystemExit(main())
